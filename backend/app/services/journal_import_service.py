"""Загрузка выгрузки электронного журнала преподавателя (по одной дисциплине и группе): оценки и «н» по часам.

Формат (проверен по образцу колледжа, файл делает Apache POI):
  * лист состоит из страниц, поставленных друг под другом: строка с «№» / «Студент» — шапка страницы;
    в ней объединённые ячейки с названием месяца, строкой ниже — число месяца над каждым столбцом
    (столбец = один академический час, у соседних часов одного дня число одно; у объединённого заголовка
    число стоит только в первой ячейке — для остальных берётся оно);
  * ниже — студенты (номер, ФИО), справа от дат — столбцы «Дата / Тема / Домашнее задание» и подпись
    «код дисциплины, группы, название», «Преподаватель: … период»;
  * в ячейке — пусто (присутствовал, оценки нет), «н» или оценка 2–5.
Число месяца без года: год берётся из периода в подписи (на стыке лет — по периоду), иначе из учебного года.

Данные кладутся в `lesson_marks` отдельно от дневных отметок куратора: «н» из журнала преподавателя не меняет
посещаемость (причину пропуска знает куратор), а даёт сверку и успеваемость. ФИО преподавателя в платформу
не сохраняется. Повторная загрузка того же файла ничего не дублирует; для дат, которые есть в файле,
журнал считается главным (стёртая преподавателем оценка стирается и здесь)."""
import datetime
import io
import re
from collections import Counter
from dataclasses import dataclass, field

import openpyxl
from sqlalchemy.orm import Session

from app.core import policies
from app.core.xlsx import check_zip_safety
from app.models import JournalLesson, LessonMark, Student, StudentStatus, StudyGroup, User
from app.services.audit_service import log_action

MAX_LISTED = 200
FIRST_DATE_COLUMN = 3  # C: после «№» и «Студент»
MONTHS = {
    "янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "мая": 5, "июн": 6,
    "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12,
}
_ABSENT_VALUES = {"н", "h"}  # латинская «h» — частая подмена кириллической «н» при вводе
_GRADES = {"2", "3", "4", "5"}
_PERIOD = re.compile(r"(\d{2}\.\d{2}\.\d{4})\s*[-–—]\s*(\d{2}\.\d{2}\.\d{4})")
_LESSON_DATE = re.compile(r"^(\d{1,2})\.(\d{2})$")
NO_HOMEWORK = "не задано"
_GROUP_TOKEN = re.compile(r"[A-Za-zА-ЯЁа-яё]{1,5}\d{3,5}[A-Za-zА-ЯЁа-яё]?(?:\.\d+)?(?![\w.])")


class ImportFileError(ValueError):
    """Файл нельзя прочитать как выгрузку журнала (текст показывается пользователю)."""


@dataclass
class Cell:
    date: datetime.date
    slot: int
    grade: int | None
    absent: bool


@dataclass
class JournalStudent:
    name: str
    cells: list[Cell] = field(default_factory=list)


@dataclass
class Lesson:
    date: datetime.date
    slot: int
    topic: str | None = None
    homework: str | None = None


@dataclass
class ParsedJournal:
    subject_code: str = ""
    subject_name: str = ""
    group_tokens: list[str] = field(default_factory=list)
    period: tuple[datetime.date, datetime.date] | None = None
    students: list[JournalStudent] = field(default_factory=list)
    lessons: list[Lesson] = field(default_factory=list)  # по одному на столбец, тема — из списка уроков справа
    dates: set[datetime.date] = field(default_factory=set)  # все даты-столбцы файла, в том числе без отметок
    unknown_values: Counter = field(default_factory=Counter)
    warnings: list[str] = field(default_factory=list)


def norm(value: str | None) -> str:
    """Для сопоставления ФИО: без регистра, ё = е, лишние пробелы убраны."""
    return " ".join((value or "").replace("ё", "е").replace("Ё", "Е").split()).casefold()


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _anchor(ws, merged: dict[tuple[int, int], tuple[int, int]], row: int, col: int):
    """Значение ячейки; у объединённой — значение её первой (левой верхней) ячейки."""
    row, col = merged.get((row, col), (row, col))
    return ws.cell(row, col).value


def _merged_map(ws) -> dict[tuple[int, int], tuple[int, int]]:
    result: dict[tuple[int, int], tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                result[(r, c)] = (rng.min_row, rng.min_col)
    return result


def _parse_date_text(text: str) -> datetime.date:
    return datetime.datetime.strptime(text, "%d.%m.%Y").date()


def _academic_year_start(today: datetime.date) -> int:
    return today.year if today.month >= 9 else today.year - 1


def _resolve_year(month: int, day: int, period, today: datetime.date) -> datetime.date | None:
    """Дата по числу и месяцу: год выбирается так, чтобы дата попала в период журнала."""
    years = [period[0].year, period[1].year] if period else []
    if not years:
        start = _academic_year_start(today)
        years = [start, start + 1]
    for year in dict.fromkeys(years):
        try:
            candidate = datetime.date(year, month, day)
        except ValueError:
            continue
        if period is None or period[0] <= candidate <= period[1]:
            return candidate
    if period is None:
        return None
    # вне периода: оставляем по учебному году (сентябрь–декабрь — год начала периода)
    year = period[0].year if month >= 9 else period[1].year
    try:
        return datetime.date(year, month, day)
    except ValueError:
        return None


def _find_headers(ws) -> list[int]:
    rows = []
    for r in range(1, ws.max_row + 1):
        if _text(ws.cell(r, 1).value) == "№" and norm(_text(ws.cell(r, 2).value)).startswith("студент"):
            rows.append(r)
    return rows


def _footer(ws, start: int, end: int, date_col: int) -> tuple[str | None, str | None]:
    """Строки подписи в столбце «Дата»: (строка «код … группы … название», строка «Преподаватель: … период»)."""
    texts = [(r, _text(ws.cell(r, date_col).value)) for r in range(start, end + 1)]
    texts = [(r, t) for r, t in texts if t]
    for i, (_, text) in enumerate(texts):
        if text.casefold().startswith("преподаватель"):
            subject_line = texts[i - 1][1] if i > 0 else None
            return subject_line, text
    return None, None


def _parse_subject_line(line: str) -> tuple[str, str, list[str]]:
    parts = line.split(maxsplit=1)
    code = parts[0].strip(" ,")
    rest = parts[1] if len(parts) > 1 else ""
    if "," in rest:
        name = rest.rsplit(",", 1)[1].strip()
    else:
        name = rest.strip()
    tokens = _GROUP_TOKEN.findall(rest)
    if not name:
        name = rest
    return code, name, tokens


def _attach_topics(ws, merged, parsed: ParsedJournal, top: int, end: int, date_col: int, period, today) -> None:
    """Темы и домашние задания из списка уроков справа от дат: строка списка = один час, порядок внутри
    дня совпадает с порядком столбцов. Лишние столбцы дня (в образце бывает 5 столбцов на 4 строки списка)
    остаются без темы."""
    by_day: dict[datetime.date, list[tuple[str | None, str | None]]] = {}
    for row in range(top + 2, end + 1):
        if merged.get((row, date_col), (row, date_col))[0] != row:
            continue  # вторая половина объединённой пары строк
        m = _LESSON_DATE.match(_text(ws.cell(row, date_col).value))
        if not m:
            continue
        day = _resolve_year(int(m.group(2)), int(m.group(1)), period, today)
        if day is None:
            continue
        topic = _text(ws.cell(row, date_col + 1).value) or None
        homework = _text(ws.cell(row, date_col + 2).value) or None
        if homework and homework.casefold() == NO_HOMEWORK:
            homework = None
        by_day.setdefault(day, []).append((topic, homework))
    for lesson in parsed.lessons:
        items = by_day.get(lesson.date, [])
        if lesson.topic is None and lesson.slot <= len(items):
            lesson.topic, lesson.homework = items[lesson.slot - 1]


def parse_workbook(content: bytes, today: datetime.date) -> ParsedJournal:
    try:
        check_zip_safety(content)
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:  # noqa: BLE001 — любой сбой чтения = «не тот файл»
        raise ImportFileError("Не удалось прочитать файл: нужна выгрузка электронного журнала в формате .xlsx") from exc
    ws = wb.worksheets[0]
    headers = _find_headers(ws)
    if not headers:
        raise ImportFileError("Не найдена шапка журнала («№», «Студент»): это не выгрузка электронного журнала")
    merged = _merged_map(ws)
    parsed = ParsedJournal()
    bounds = headers + [ws.max_row + 1]

    # Подпись есть на каждой странице, и период в ней свой (стр. 1 — 01.09–22.09, стр. 2 — 29.09–10.11):
    # период журнала — от самого раннего начала до самого позднего конца.
    periods: list[tuple[datetime.date, datetime.date]] = []
    for idx, h in enumerate(headers):
        date_col = next((c for c in range(FIRST_DATE_COLUMN, ws.max_column + 1) if norm(_text(ws.cell(h, c).value)) == "дата"), None)
        if date_col is None:
            continue
        subject_line, teacher_line = _footer(ws, h, bounds[idx + 1] - 1, date_col)
        if subject_line and not parsed.subject_code:
            parsed.subject_code, parsed.subject_name, parsed.group_tokens = _parse_subject_line(subject_line)
        m = _PERIOD.search(teacher_line or "")
        if m:
            try:
                periods.append((_parse_date_text(m.group(1)), _parse_date_text(m.group(2))))
            except ValueError:
                pass
    period = (min(p[0] for p in periods), max(p[1] for p in periods)) if periods else None
    if not parsed.subject_code:
        raise ImportFileError("Не найдена подпись журнала с кодом дисциплины («МДК.01.01 …»): без неё непонятно, к какому предмету оценки")
    parsed.period = period
    if period is None:
        parsed.warnings.append("В подписи не найден период журнала — год дат определён по учебному году")

    by_name: dict[str, JournalStudent] = {}
    slots_by_date: Counter = Counter()
    for idx, h in enumerate(headers):
        end = bounds[idx + 1] - 1
        date_col = next((c for c in range(FIRST_DATE_COLUMN, ws.max_column + 1) if norm(_text(ws.cell(h, c).value)) == "дата"), ws.max_column + 1)
        columns: list[tuple[int, datetime.date, int]] = []
        for col in range(FIRST_DATE_COLUMN, date_col):
            month_text = norm(_text(_anchor(ws, merged, h, col)))[:3]
            day_text = _text(_anchor(ws, merged, h + 1, col))
            if not day_text:
                continue
            month = MONTHS.get(month_text)
            if month is None or not day_text.isdigit():
                parsed.warnings.append(f"Строка {h}, столбец {col}: не распознаны месяц или число — столбец пропущен")
                continue
            day = _resolve_year(month, int(day_text), period, today)
            if day is None:
                parsed.warnings.append(f"Строка {h + 1}, столбец {col}: такой даты не бывает ({day_text}.{month:02d}) — столбец пропущен")
                continue
            if period and not (period[0] <= day <= period[1]) and day not in parsed.dates:
                parsed.warnings.append(f"Дата {day:%d.%m.%Y} выходит за период журнала")
            slots_by_date[day] += 1
            columns.append((col, day, slots_by_date[day]))
            parsed.dates.add(day)
            parsed.lessons.append(Lesson(day, slots_by_date[day]))
        _attach_topics(ws, merged, parsed, h, end, date_col, period, today)
        row = h + 2
        while row <= end and _text(ws.cell(row, 1).value).isdigit() and _text(ws.cell(row, 2).value):
            name = _text(ws.cell(row, 2).value)
            student = by_name.setdefault(norm(name), JournalStudent(name=name))
            for col, day, slot in columns:
                raw = _text(ws.cell(row, col).value)
                if not raw:
                    continue
                key = norm(raw)
                if key in _GRADES:
                    student.cells.append(Cell(day, slot, int(key), False))
                elif key in _ABSENT_VALUES:
                    student.cells.append(Cell(day, slot, None, True))
                else:
                    parsed.unknown_values[raw] += 1
            row += 1
    parsed.students = list(by_name.values())
    if not parsed.students:
        raise ImportFileError("В журнале нет ни одного студента")
    return parsed


# --- загрузка в базу ---------------------------------------------------------------------------------

@dataclass
class Report:
    counts: dict[str, int] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    not_found: list[str] = field(default_factory=list)
    ambiguous: list[str] = field(default_factory=list)
    out_of_scope: list[str] = field(default_factory=list)
    subject_code: str = ""
    subject_name: str = ""
    period: tuple[datetime.date, datetime.date] | None = None
    blocked: str | None = None


def code_in_tokens(code: str | None, tokens: set[str]) -> bool:
    """Код группы из базы есть среди кодов подписи журнала (с точностью до «.9» — номера подгруппы)."""
    c = norm(code)
    return bool(c) and (c in tokens or c.split(".")[0] in tokens or any(t.split(".")[0] == c for t in tokens))


def group_key(tokens: list[str]) -> str:
    """Группы журнала одной строкой для поиска «есть ли у журнала группа X»: ,са1112,са2112.9,"""
    return "," + ",".join(sorted({norm(t) for t in tokens})) + ","


def _store_lessons(db: Session, parsed: ParsedJournal, counts: Counter) -> None:
    """Столбцы журнала (дата, час, тема, домашнее задание): по ним рисуется журнал в платформе, в том числе часы,
    где ни у кого нет отметок. Журнал главный: часы, которых в новом файле нет, убираются."""
    key = group_key(parsed.group_tokens)
    existing = {
        (l.date, l.slot): l
        for l in db.query(JournalLesson).filter(
            JournalLesson.subject_code == parsed.subject_code, JournalLesson.group_key == key,
            JournalLesson.date.in_(parsed.dates),
        )
    }
    keep = set()
    for lesson in parsed.lessons:
        keep.add((lesson.date, lesson.slot))
        row = existing.get((lesson.date, lesson.slot))
        if row is None:
            db.add(JournalLesson(
                subject_code=parsed.subject_code, subject_name=parsed.subject_name, group_key=key,
                date=lesson.date, slot=lesson.slot, topic=lesson.topic, homework=lesson.homework,
            ))
            counts["lessons_new"] += 1
        elif (row.topic, row.homework, row.subject_name) != (lesson.topic, lesson.homework, parsed.subject_name):
            row.topic, row.homework, row.subject_name = lesson.topic, lesson.homework, parsed.subject_name
            counts["lessons_updated"] += 1
    for k, row in existing.items():
        if k not in keep:
            db.delete(row)


def process(db: Session, parsed: ParsedJournal, user: User, today: datetime.date) -> Report:
    """Разбор + запись (вызывающий коммитит или откатывает: предпросмотр — rollback)."""
    report = Report(subject_code=parsed.subject_code, subject_name=parsed.subject_name, period=parsed.period)
    report.warnings.extend(parsed.warnings)
    for value, count in parsed.unknown_values.most_common(10):
        report.warnings.append(f"Значение «{value}» встретилось {count} раз(а) — не оценка и не «н», пропущено")

    scope = policies.scope_of(user, "edit")
    allowed_group_ids = {g for (g,) in db.query(StudyGroup.id).filter(scope.group_filter(db, today))}
    students: dict[str, list[Student]] = {}
    for s in db.query(Student).filter(Student.status != StudentStatus.EXPELLED):
        students.setdefault(norm(s.full_name), []).append(s)

    group_codes = {g.id: g.code for g in db.query(StudyGroup)}
    file_tokens = {norm(t) for t in parsed.group_tokens}
    matched: list[tuple[JournalStudent, Student]] = []
    for js in parsed.students:
        candidates = students.get(norm(js.name), [])
        if len(candidates) > 1:
            # Однофамильцы с одинаковым ФИО: сначала по группам из подписи журнала, потом по зоне пользователя.
            for keep in (lambda s: code_in_tokens(group_codes.get(s.study_group_id), file_tokens), lambda s: s.study_group_id in allowed_group_ids):
                narrowed = [s for s in candidates if keep(s)]
                if narrowed:
                    candidates = narrowed
                if len(candidates) == 1:
                    break
        if not candidates:
            report.not_found.append(js.name)
        elif len(candidates) > 1:
            report.ambiguous.append(js.name)
        elif candidates[0].study_group_id not in allowed_group_ids:
            report.out_of_scope.append(js.name)
        else:
            matched.append((js, candidates[0]))

    foreign = [s for _, s in matched if file_tokens and not code_in_tokens(group_codes.get(s.study_group_id), file_tokens)]
    if foreign:
        report.warnings.append(f"{len(foreign)} студент(ов) числятся в группе, которой нет в подписи журнала ({', '.join(parsed.group_tokens)})")
    if report.not_found:
        report.warnings.append(f"Не найдено в платформе: {len(report.not_found)} (проверьте написание ФИО)")
    if report.out_of_scope:
        report.warnings.append(f"Не из ваших групп, пропущено: {len(report.out_of_scope)}")
    if report.ambiguous:
        report.warnings.append(f"Однофамильцы с одинаковым ФИО, пропущено: {len(report.ambiguous)}")

    counts = Counter(students_in_file=len(parsed.students), students_matched=len(matched))
    report.counts = counts  # type: ignore[assignment]
    if not matched:
        report.blocked = "Ни один студент журнала не найден среди ваших студентов: ничего не загружено"
        report.counts = dict(counts)
        return report

    ids = [s.id for _, s in matched]
    existing = {
        (m.student_id, m.date, m.slot): m
        for m in db.query(LessonMark).filter(LessonMark.subject_code == parsed.subject_code, LessonMark.student_id.in_(ids), LessonMark.date.in_(parsed.dates))
    }
    seen: set[tuple[int, datetime.date, int]] = set()
    now_user = user.id
    for js, student in matched:
        for cell in js.cells:
            key = (student.id, cell.date, cell.slot)
            seen.add(key)
            row = existing.get(key)
            if row is None:
                db.add(LessonMark(
                    student_id=student.id, subject_code=parsed.subject_code, subject_name=parsed.subject_name,
                    date=cell.date, slot=cell.slot, grade=cell.grade, absent=cell.absent, imported_by_user_id=now_user,
                ))
                counts["grades_new" if cell.grade is not None else "absences_new"] += 1
            elif row.grade != cell.grade or row.absent != cell.absent:
                row.grade, row.absent, row.subject_name, row.imported_by_user_id = cell.grade, cell.absent, parsed.subject_name, now_user
                counts["updated"] += 1
            else:
                counts["unchanged"] += 1
    for key, row in existing.items():
        if key not in seen:
            db.delete(row)
            counts["removed"] += 1
    _store_lessons(db, parsed, counts)
    report.counts = dict(counts)
    for name in ("grades_new", "absences_new", "updated", "unchanged", "removed", "lessons_new", "lessons_updated"):
        report.counts.setdefault(name, 0)
    db.flush()
    log_action(
        db, user, "journal.import", "subject", parsed.subject_code,
        new_value=f"студентов={len(matched)} оценок+={report.counts['grades_new']} н+={report.counts['absences_new']} "
                  f"изменено={report.counts['updated']} удалено={report.counts['removed']}",
    )
    return report
