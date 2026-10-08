"""Единая загрузка контингента из одного Excel-шаблона: группы, студенты, кураторы.

Шаблон — три листа: «Группы», «Студенты», «Кураторы» (+ «Инструкция», её загрузка игнорирует). Его можно скачать
с текущими данными платформы, поправить и загрузить обратно: что есть в таблице — создаётся или обновляется,
остальное остаётся как было. Листы необязательны: можно загрузить только кураторов или только одну группу.

Принципы (те же, что у консольных загрузок, но с предпросмотром):
  * сначала предпросмотр — те же вычисления, но результат откатывается; запись — повторный прогон тех же правил;
  * любая ошибка в файле блокирует запись целиком (частично применённый контингент хуже непримённого);
  * пустая ячейка = «не менять» (у нового — значение по умолчанию);
  * студентов, которых в файле нет, по умолчанию не трогаем; «выбывшими» их можно отметить только для групп,
    которые в файле есть, и только после подтверждения, если их слишком много (защита от неполного файла);
  * никогда ничего не удаляем — выбытие это статус «отчислен», история и отметки сохраняются;
  * временные пароли выдаются по желанию и только при записи; в журнал аудита пароли не попадают.
Работает без зависимостей от `app.api` — им пользуются и экран администратора, и консольная команда.
"""
import datetime
import io
import re
from collections import defaultdict
from dataclasses import dataclass, field
from itertools import islice

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.core.time import today_local
from app.core.xlsx import UnsafeArchiveError, append_row, check_zip_safety
from app.models import (
    AssignmentRole, CuratorAssignment, Department, Role, RoleCode, Student, StudentStatus, StudyGroup, User,
)
from app.services import group_membership_service
from app.services.password_service import generate_temporary_password

MAX_ROWS = 5000
MAX_SCANNED_ROWS = MAX_ROWS + 2000
GUARD_MIN_ABSENT = 10          # меньше этого числа «выбывших» подтверждения не требуем
GUARD_PERCENT = 10.0           # доля студентов затронутых групп, выше которой нужно подтверждение
MAX_CHANGES_LISTED = 400       # сколько строк подробного списка отдаём на экран (счётчики — всегда полные)

GROUPS, STUDENTS, CURATORS, HELP = "Группы", "Студенты", "Кураторы", "Инструкция"

GROUP_COLUMNS = [("code", "Код группы"), ("course", "Курс"), ("department", "Отделение"),
                 ("funding", "Финансирование (бюджет/договор)"), ("active", "Группа действует (да/нет)")]
STUDENT_COLUMNS = [("last_name", "Фамилия"), ("first_name", "Имя"), ("middle_name", "Отчество"), ("group", "Группа"),
                   ("status", "Статус (обучается / академ. отпуск / отчислен)"), ("enrolled_at", "Дата зачисления")]
CURATOR_COLUMNS = [("full_name", "ФИО куратора"), ("group", "Группа"),
                   ("role", "Роль (куратор / заместитель)"), ("start", "С какой даты")]
REQUIRED = {GROUPS: {"code"}, STUDENTS: {"last_name", "first_name", "group"}, CURATORS: {"full_name", "group"}}

_TRUE = {"да", "д", "yes", "y", "true", "1", "+"}
_FALSE = {"нет", "н", "no", "n", "false", "0", "-"}
_FUNDING = {"бюджет": "budget", "budget": "budget", "договор": "contract", "контракт": "contract", "contract": "contract"}
_STATUS = {
    "обучается": StudentStatus.STUDYING, "учится": StudentStatus.STUDYING, "studying": StudentStatus.STUDYING,
    "академ. отпуск": StudentStatus.ACADEMIC_LEAVE, "академический отпуск": StudentStatus.ACADEMIC_LEAVE,
    "в академическом отпуске": StudentStatus.ACADEMIC_LEAVE, "академ": StudentStatus.ACADEMIC_LEAVE,
    "academic_leave": StudentStatus.ACADEMIC_LEAVE,
    "отчислен": StudentStatus.EXPELLED, "выбыл": StudentStatus.EXPELLED, "expelled": StudentStatus.EXPELLED,
}
_STATUS_TEXT = {StudentStatus.STUDYING: "обучается", StudentStatus.ACADEMIC_LEAVE: "академ. отпуск", StudentStatus.EXPELLED: "отчислен"}
_ROLE = {"куратор": AssignmentRole.CURATOR, "curator": AssignmentRole.CURATOR,
         "заместитель": AssignmentRole.DEPUTY, "зам": AssignmentRole.DEPUTY, "зам. куратора": AssignmentRole.DEPUTY,
         "заместитель куратора": AssignmentRole.DEPUTY, "deputy": AssignmentRole.DEPUTY}
_YEAR_SUFFIX = re.compile(r"-\d{2}")

_TRANSLIT = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh", "з": "z", "и": "i",
    "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t",
    "у": "u", "ф": "f", "х": "kh", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "shch", "ъ": "", "ы": "y",
    "ь": "", "э": "e", "ю": "yu", "я": "ya",
}


class ImportFileError(ValueError):
    """Файл нельзя разобрать целиком (не xlsx, нет нужных листов, слишком большой)."""


# ---------------------------------------------------------------------------------------------------------------
# Разбор файла
# ---------------------------------------------------------------------------------------------------------------

@dataclass
class RowError:
    sheet: str
    row: int
    message: str


@dataclass
class GroupRow:
    row: int
    code: str
    course: int | None
    department: str | None
    funding: str | None
    active: bool | None


@dataclass
class StudentRow:
    row: int
    last_name: str
    first_name: str
    middle_name: str | None
    group: str
    status: StudentStatus | None
    enrolled_at: datetime.date | None


@dataclass
class CuratorRow:
    row: int
    full_name: str
    group: str
    role: AssignmentRole
    start: datetime.date | None


@dataclass
class ParsedFile:
    groups: list[GroupRow] = field(default_factory=list)
    students: list[StudentRow] = field(default_factory=list)
    curators: list[CuratorRow] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    sheets: list[str] = field(default_factory=list)  # какие из рабочих листов найдены


def norm(value: str | None) -> str:
    """Для сопоставления: без регистра, ё = е, лишние пробелы убраны."""
    return " ".join((value or "").replace("ё", "е").replace("Ё", "Е").split()).casefold()


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _date(value, what: str) -> datetime.date | None:
    if value is None or _text(value) == "":
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = _text(value)
    for pattern in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y", "%d/%m/%Y"):
        try:
            return datetime.datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError(f"{what}: «{text}» — нужна дата вида 01.09.2026")


def _columns(sheet_name: str, header: tuple, spec: list[tuple[str, str]]) -> dict[str, int]:
    """Колонки ищутся по заголовку (по началу названия, без регистра) — порядок и лишние колонки не важны."""
    titles = [norm(_text(c)) for c in header]
    found: dict[str, int] = {}
    for key, title in spec:
        prefix = norm(title).split(" (")[0]
        for index, existing in enumerate(titles):
            if existing.startswith(prefix) and index not in found.values():
                found[key] = index
                break
    missing = [title for key, title in spec if key in REQUIRED[sheet_name] and key not in found]
    if missing:
        raise ImportFileError(f"Лист «{sheet_name}»: не найдены обязательные колонки — {', '.join(missing)}")
    return found


def parse_workbook(content: bytes) -> ParsedFile:
    try:
        check_zip_safety(content)
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except UnsafeArchiveError as exc:
        raise ImportFileError(f"Файл не принят: {exc}") from exc
    except Exception as exc:  # noqa: BLE001 — любой сбой разбора xlsx для пользователя одинаков
        raise ImportFileError("Не удалось прочитать файл — нужен Excel (.xlsx), сохранённый из шаблона платформы") from exc

    parsed = ParsedFile()
    names = {norm(n): n for n in workbook.sheetnames}
    for sheet_name, spec in ((GROUPS, GROUP_COLUMNS), (STUDENTS, STUDENT_COLUMNS), (CURATORS, CURATOR_COLUMNS)):
        if norm(sheet_name) not in names:
            continue
        parsed.sheets.append(sheet_name)
        rows = list(islice(workbook[names[norm(sheet_name)]].iter_rows(values_only=True), MAX_SCANNED_ROWS + 1))
        if not rows:
            continue
        columns = _columns(sheet_name, rows[0], spec)
        data = [(i, r) for i, r in enumerate(rows[1:], start=2) if any(_text(c) for c in r)]
        if len(data) > MAX_ROWS:
            raise ImportFileError(f"Лист «{sheet_name}»: больше {MAX_ROWS} строк — разбейте файл на части")
        for number, row in data:
            def cell(key: str, row=row):
                index = columns.get(key)
                return row[index] if index is not None and index < len(row) else None
            try:
                _parse_row(parsed, sheet_name, number, cell)
            except ValueError as exc:
                parsed.errors.append(RowError(sheet_name, number, str(exc)))
    workbook.close()
    if not parsed.sheets:
        raise ImportFileError(f"В файле нет ни одного из листов «{GROUPS}», «{STUDENTS}», «{CURATORS}» — скачайте шаблон заново")
    return parsed


def _parse_row(parsed: ParsedFile, sheet: str, number: int, cell) -> None:
    if sheet == GROUPS:
        code = _text(cell("code"))
        if not code:
            raise ValueError("не указан код группы")
        if len(code) > 32:
            raise ValueError("код группы длиннее 32 знаков")
        course_text = _text(cell("course"))
        course = None
        if course_text:
            match = re.match(r"\d+", course_text)
            if match is None or not 1 <= int(match.group()) <= 6:
                raise ValueError(f"курс «{course_text}» — нужно число от 1 до 6")
            course = int(match.group())
        funding_text = norm(_text(cell("funding")))
        if funding_text and funding_text not in _FUNDING:
            raise ValueError(f"финансирование «{_text(cell('funding'))}» — нужно «бюджет» или «договор»")
        active_text = norm(_text(cell("active")))
        if active_text and active_text not in _TRUE | _FALSE:
            raise ValueError(f"«Группа действует»: «{_text(cell('active'))}» — нужно «да» или «нет»")
        parsed.groups.append(GroupRow(
            row=number, code=code, course=course, department=_text(cell("department")) or None,
            funding=_FUNDING.get(funding_text), active=None if not active_text else active_text in _TRUE,
        ))
    elif sheet == STUDENTS:
        last, first = _text(cell("last_name")), _text(cell("first_name"))
        if not last or not first or not _text(cell("group")):
            raise ValueError("нужны фамилия, имя и группа")
        if max(len(last), len(first), len(_text(cell("middle_name")))) > 128:
            raise ValueError("фамилия, имя или отчество длиннее 128 знаков")
        status_text = norm(_text(cell("status")))
        if status_text and status_text not in _STATUS:
            raise ValueError(f"статус «{_text(cell('status'))}» — нужно «обучается», «академ. отпуск» или «отчислен»")
        parsed.students.append(StudentRow(
            row=number, last_name=last, first_name=first, middle_name=_text(cell("middle_name")) or None,
            group=_text(cell("group")), status=_STATUS.get(status_text), enrolled_at=_date(cell("enrolled_at"), "дата зачисления"),
        ))
    else:
        full_name = _text(cell("full_name"))
        if len(full_name.split()) < 2 or not _text(cell("group")):
            raise ValueError("нужны ФИО куратора (фамилия и имя) и группа")
        role_text = norm(_text(cell("role")))
        if role_text and role_text not in _ROLE:
            raise ValueError(f"роль «{_text(cell('role'))}» — нужно «куратор» или «заместитель»")
        parsed.curators.append(CuratorRow(
            row=number, full_name=full_name, group=_text(cell("group")),
            role=_ROLE.get(role_text, AssignmentRole.CURATOR), start=_date(cell("start"), "дата начала"),
        ))


# ---------------------------------------------------------------------------------------------------------------
# Применение (общее для предпросмотра и записи)
# ---------------------------------------------------------------------------------------------------------------

@dataclass
class Options:
    absent_students: str = "keep"          # keep — не трогать; expel — отметить «отчислен» (только в группах из файла)
    replace_curators: bool = False         # назначить нового куратора, закрыв действующего
    create_departments: bool = False       # заводить отделения, которых нет на платформе
    confirm_large: bool = False            # подтверждение массового выбытия
    issue_passwords: bool = False          # выдать временные пароли кураторам без пароля
    default_enrolled_at: datetime.date | None = None
    today: datetime.date | None = None


@dataclass
class Change:
    sheet: str
    action: str
    label: str
    detail: str = ""


@dataclass
class Credential:
    full_name: str
    department: str
    username: str
    password: str


@dataclass
class Report:
    counts: dict[str, int] = field(default_factory=dict)
    changes: list[Change] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    needs_confirmation: str | None = None
    credentials: list[Credential] = field(default_factory=list)

    def count(self, name: str, n: int = 1) -> None:
        self.counts[name] = self.counts.get(name, 0) + n

    def add(self, sheet: str, action: str, label: str, detail: str = "") -> None:
        self.changes.append(Change(sheet, action, label, detail))

    @property
    def blocked(self) -> bool:
        return bool(self.errors) or self.needs_confirmation is not None


def _translit(text: str) -> str:
    return "".join(_TRANSLIT.get(ch, ch) for ch in text.lower() if ch.isalpha() or ch in _TRANSLIT)


def base_username(full_name: str) -> str:
    parts = full_name.split()
    return f"{_translit(parts[0])}.{_translit(parts[1])[:1]}"


def _resolve_group(code: str, groups: dict[str, StudyGroup]) -> StudyGroup | None:
    """Группа по коду: точно (без регистра) или «код-ГГ» (год набора в реестре), если такая одна."""
    exact = groups.get(norm(code))
    if exact is not None:
        return exact
    found = [g for c, g in groups.items() if c.startswith(norm(code)) and _YEAR_SUFFIX.fullmatch(c[len(norm(code)):])]
    return found[0] if len(found) == 1 else None


def process(db: Session, parsed: ParsedFile, options: Options, user: User | None = None) -> Report:
    """Выполняет правила над сессией (flush, без commit) и возвращает отчёт. Предпросмотр = process + rollback."""
    report = Report(errors=list(parsed.errors))
    today = options.today or today_local()
    enrolled_default = options.default_enrolled_at or today

    departments = {norm(d.name): d for d in db.query(Department).all()}
    groups: dict[str, StudyGroup] = {norm(g.code): g for g in db.query(StudyGroup).all()}

    # ---- группы ----
    seen_codes: set[str] = set()
    for r in parsed.groups:
        if norm(r.code) in seen_codes:
            report.errors.append(RowError(GROUPS, r.row, f"группа {r.code} указана в файле дважды"))
            continue
        seen_codes.add(norm(r.code))
        department = None
        if r.department:
            department = departments.get(norm(r.department))
            if department is None:
                if options.create_departments:
                    department = Department(name=r.department)
                    db.add(department)
                    db.flush()
                    departments[norm(r.department)] = department
                    report.count("отделений создано")
                    report.add(GROUPS, "создано", f"Отделение {r.department}")
                else:
                    report.errors.append(RowError(GROUPS, r.row, f"отделение «{r.department}» не найдено на платформе "
                                                  "(опечатка? или включите «создавать новые отделения»)"))
                    continue
        group = groups.get(norm(r.code))
        if group is None:
            if department is None or r.course is None:
                report.errors.append(RowError(GROUPS, r.row, f"новая группа {r.code}: нужны курс и отделение"))
                continue
            group = StudyGroup(code=r.code, course=r.course, department_id=department.id,
                               funding=r.funding, is_active=True if r.active is None else r.active)
            db.add(group)
            db.flush()
            groups[norm(r.code)] = group
            report.count("групп создано")
            report.add(GROUPS, "создана", f"Группа {r.code}", f"{r.course} курс, {department.name}")
            continue
        changed = []
        if r.course is not None and r.course != group.course:
            changed.append(f"курс {group.course} → {r.course}")
            group.course = r.course
        if department is not None and department.id != group.department_id:
            changed.append(f"отделение → {department.name}")
            group.department_id = department.id
        if r.funding is not None and r.funding != (group.funding or "budget"):
            changed.append(f"финансирование → {'договор' if r.funding == 'contract' else 'бюджет'}")
            group.funding = r.funding
        if r.active is not None and r.active != group.is_active:
            changed.append("возвращена в работу" if r.active else "скрыта")
            group.is_active = r.active
        if changed:
            report.count("групп изменено")
            report.add(GROUPS, "изменена", f"Группа {group.code}", "; ".join(changed))
    db.flush()

    # ---- студенты ----
    all_students = db.query(Student).all()
    by_group: dict[int, list[Student]] = defaultdict(list)
    for s in all_students:
        by_group[s.study_group_id].append(s)

    def fio(s: Student) -> tuple[str, str, str]:
        return (norm(s.last_name), norm(s.first_name), norm(s.middle_name))

    rows_ok: list[tuple[StudentRow, StudyGroup]] = []
    seen_students: set[tuple[int, str, str, str]] = set()
    for r in parsed.students:
        group = _resolve_group(r.group, groups)
        if group is None:
            report.errors.append(RowError(STUDENTS, r.row, f"группа «{r.group}» не найдена (добавьте её на лист «{GROUPS}» или поправьте код)"))
            continue
        key = (group.id, norm(r.last_name), norm(r.first_name), norm(r.middle_name))
        if key in seen_students:
            report.errors.append(RowError(STUDENTS, r.row, f"{r.last_name} {r.first_name} в группе {group.code} указан в файле дважды"))
            continue
        seen_students.add(key)
        rows_ok.append((r, group))

    matched: dict[int, tuple[StudentRow, StudyGroup]] = {}
    unmatched: list[tuple[StudentRow, StudyGroup]] = []
    for r, group in rows_ok:
        want = (norm(r.last_name), norm(r.first_name), norm(r.middle_name))
        pool = [s for s in by_group[group.id] if s.id not in matched
                and fio(s)[:2] == want[:2] and (not want[2] or fio(s)[2] == want[2])]
        if pool:
            pool.sort(key=lambda s: (s.status == StudentStatus.EXPELLED, s.id))
            matched[pool[0].id] = (r, group)
        else:
            unmatched.append((r, group))

    free = [s for s in all_students if s.id not in matched]
    for r, group in unmatched:
        want = (norm(r.last_name), norm(r.first_name), norm(r.middle_name))
        pool = [s for s in free if s.id not in matched and fio(s)[:2] == want[:2] and (not want[2] or fio(s)[2] == want[2])]
        if len(pool) == 1:
            matched[pool[0].id] = (r, group)
            continue
        status = r.status or StudentStatus.STUDYING
        student = Student(last_name=r.last_name, first_name=r.first_name, middle_name=r.middle_name,
                          study_group_id=group.id, status=status, enrolled_at=r.enrolled_at or enrolled_default,
                          left_at=today if status == StudentStatus.EXPELLED else None)
        db.add(student)
        db.flush()
        group_membership_service.create_initial_membership(db, student)
        report.count("студентов создано")
        report.add(STUDENTS, "создан", student.full_name, f"группа {group.code}")

    by_id = {s.id: s for s in all_students}
    for student_id, (r, group) in matched.items():
        student = by_id[student_id]
        changes = []
        if student.study_group_id != group.id:
            old_code = student.study_group.code if student.study_group else str(student.study_group_id)
            group_membership_service.transfer_student(db, student, group.id, today, user)
            student.study_group_id = group.id
            changes.append(f"перевод {old_code} → {group.code}")
            report.count("студентов переведено")
        if r.status is not None and r.status != student.status:
            changes.append(f"статус: {_STATUS_TEXT[student.status]} → {_STATUS_TEXT[r.status]}")
            student.status = r.status
            student.left_at = today if r.status == StudentStatus.EXPELLED else None
            report.count("статус студента изменён")
        elif r.status is None and student.status == StudentStatus.EXPELLED:
            # Отчисленный указан в файле без статуса — значит, он снова в списке группы.
            changes.append("возвращён из отчисленных")
            student.status = StudentStatus.STUDYING
            student.left_at = None
            report.count("студентов возвращено из отчисленных")
        if changes:
            report.add(STUDENTS, "изменён", student.full_name, "; ".join(changes))
    db.flush()

    # Студенты затронутых групп, которых в файле нет.
    if parsed.students:
        listed_groups = {g.id for _, g in rows_ok}
        absent = [s for gid in listed_groups for s in by_group[gid]
                  if s.id not in matched and s.status != StudentStatus.EXPELLED]
        if absent:
            total = sum(len(by_group[gid]) for gid in listed_groups) or 1
            if options.absent_students == "expel":
                share = 100 * len(absent) / total
                if len(absent) > GUARD_MIN_ABSENT and share > GUARD_PERCENT and not options.confirm_large:
                    report.needs_confirmation = (
                        f"Файл отметил бы «отчислен» {len(absent)} из {total} студентов ({share:.0f}%) — похоже на неполный список. "
                        "Проверьте файл; если выбытие настоящее, подтвердите его отдельной галочкой."
                    )
                for s in absent:
                    s.status = StudentStatus.EXPELLED
                    s.left_at = today
                    report.add(STUDENTS, "отчислен", s.full_name, f"нет в файле, группа {s.study_group.code}")
                report.count("студентов отмечено отчисленными", len(absent))
            else:
                report.warnings.append(
                    f"В файле нет {len(absent)} студентов из загруженных групп — они оставлены без изменений "
                    "(чтобы отметить их отчисленными, выберите соответствующий режим)."
                )
    db.flush()

    # ---- кураторы ----
    curator_role = db.query(Role).filter(Role.code == RoleCode.CURATOR.value).one_or_none()
    deputy_role = db.query(Role).filter(Role.code == RoleCode.DEPUTY_CURATOR.value).one_or_none()
    if parsed.curators and (curator_role is None or deputy_role is None):
        raise RuntimeError("Роли curator/deputy_curator не найдены — сначала выполните scripts.seed")
    users_by_key: dict[tuple[int | None, str], list[User]] = defaultdict(list)
    for u in db.query(User).all():
        users_by_key[(u.department_id, norm(u.full_name))].append(u)
    taken = {u for (u,) in db.query(User.username).all()}
    active_by_group: dict[tuple[int, AssignmentRole], list[CuratorAssignment]] = defaultdict(list)
    for a in db.query(CuratorAssignment).filter(CuratorAssignment.end_date.is_(None)):
        active_by_group[(a.study_group_id, a.role_type)].append(a)

    touched_users: list[User] = []
    seen_assignments: set[tuple[int, str, AssignmentRole]] = set()
    for r in parsed.curators:
        group = _resolve_group(r.group, groups)
        if group is None:
            report.errors.append(RowError(CURATORS, r.row, f"группа «{r.group}» не найдена"))
            continue
        dedupe = (group.id, norm(r.full_name), r.role)
        if dedupe in seen_assignments:
            report.errors.append(RowError(CURATORS, r.row, f"{r.full_name} для группы {group.code} указан в файле дважды"))
            continue
        seen_assignments.add(dedupe)
        candidates = users_by_key[(group.department_id, norm(r.full_name))]
        if len(candidates) > 1:
            report.errors.append(RowError(CURATORS, r.row, f"в отделении несколько пользователей «{r.full_name}» — назначьте вручную в админке"))
            continue
        if candidates:
            curator = candidates[0]
        else:
            username = candidate = base_username(r.full_name)
            suffix = 1
            while username in taken:
                suffix += 1
                username = f"{candidate}{suffix}"
            taken.add(username)
            role = deputy_role if r.role == AssignmentRole.DEPUTY else curator_role
            curator = User(username=username, full_name=r.full_name, role_id=role.id,
                           department_id=group.department_id, password_hash=None)
            db.add(curator)
            db.flush()
            users_by_key[(group.department_id, norm(r.full_name))].append(curator)
            report.count("кураторов создано")
            report.add(CURATORS, "создан", r.full_name,
                       f"логин {username}, " + ("будет выдан временный пароль" if options.issue_passwords else "без пароля"))
        touched_users.append(curator)

        start = r.start or today
        active = active_by_group[(group.id, r.role)]
        if any(a.user_id == curator.id for a in active):
            continue
        if active:
            others = ", ".join(sorted({a.user.full_name for a in active}))
            if not options.replace_curators:
                report.count("назначений пропущено (уже есть другой)")
                report.warnings.append(f"Группа {group.code}: уже назначен {others} — {r.full_name} не назначен "
                                       "(включите «заменять действующих кураторов», если нужна замена).")
                continue
            for a in active:
                a.end_date = max(a.start_date, start - datetime.timedelta(days=1))
            report.count("назначений закрыто (замена)", len(active))
        db.add(CuratorAssignment(study_group_id=group.id, user_id=curator.id, role_type=r.role, start_date=start))
        report.count("назначений создано")
        report.add(CURATORS, "назначен", r.full_name, f"группа {group.code}, с {start:%d.%m.%Y}"
                   + (f", вместо {others}" if active else ""))
    db.flush()

    if options.issue_passwords:
        seen_users: set[int] = set()
        for u in touched_users:
            if u.id in seen_users or u.password_hash is not None:
                continue
            seen_users.add(u.id)
            password = generate_temporary_password()
            u.password_hash = hash_password(password)
            u.must_change_password = True
            department = db.get(Department, u.department_id) if u.department_id else None
            report.credentials.append(Credential(u.full_name, department.name if department else "", u.username, password))
        if report.credentials:
            report.count("временных паролей выдано", len(report.credentials))
    db.flush()
    order = {GROUPS: 0, STUDENTS: 1, CURATORS: 2}
    report.errors.sort(key=lambda e: (order.get(e.sheet, 9), e.row))
    return report


# ---------------------------------------------------------------------------------------------------------------
# Шаблон
# ---------------------------------------------------------------------------------------------------------------

_INSTRUCTION = [
    "Как пользоваться шаблоном",
    "",
    "1. Заполните нужные листы: «Группы», «Студенты», «Кураторы». Листы необязательны — можно загрузить только кураторов или только одну группу.",
    "2. Первая строка каждого листа — заголовки, их менять не нужно (порядок колонок не важен). Пустая ячейка означает «не менять».",
    "3. Загрузите файл в разделе «Админка → Импорт». Платформа сначала покажет, что изменится, и только после вашего подтверждения запишет.",
    "4. Если в файле есть ошибки, запись не выполняется вообще — исправьте строки из списка и загрузите файл заново.",
    "5. Повторная загрузка того же файла ничего не меняет.",
    "",
    "Лист «Группы»: код (как на платформе, например ИИ112), курс (1–6), отделение (как в админке), финансирование (бюджет/договор), группа действует (да/нет).",
    "   Для новой группы обязательны курс и отделение. Для существующей заполняйте только то, что нужно изменить.",
    "Лист «Студенты»: фамилия, имя, отчество (можно пустое), группа, статус (обучается / академ. отпуск / отчислен), дата зачисления (для новых; по умолчанию — дата загрузки).",
    "   Студент ищется по группе и ФИО (регистр и «е/ё» не важны). Нашёлся в другой группе — переводится, не нашёлся нигде — создаётся.",
    "   Студенты, которых в файле нет, по умолчанию не трогаются. Режим «отметить отчисленными» действует только для групп, которые есть в файле.",
    "Лист «Кураторы»: ФИО, группа, роль (куратор / заместитель), с какой даты (по умолчанию — дата загрузки).",
    "   Нового куратора платформа создаёт без пароля; временные пароли выдаются галочкой «Выдать временные пароли новым кураторам».",
    "   Если у группы уже есть другой куратор, назначение пропускается; галочка «заменять действующих кураторов» закрывает прежнее назначение.",
    "",
    "Ничего не удаляется: выбытие — это статус «отчислен», история и отметки посещаемости остаются.",
]

_HEAD_FILL = PatternFill("solid", fgColor="4A2C8A")


def _sheet(workbook: Workbook, title: str, columns: list[tuple[str, str]]):
    ws = workbook.create_sheet(title)
    ws.append([header for _, header in columns])
    for i, (_, header) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = max(16, min(34, len(header) * 0.9))
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    return ws


def build_template(db: Session | None, *, with_data: bool = True) -> bytes:
    """Шаблон загрузки; с with_data — уже заполненный текущими данными платформы (удобно править и загружать обратно)."""
    workbook = Workbook()
    help_sheet = workbook.active
    help_sheet.title = HELP
    for line in _INSTRUCTION:
        help_sheet.append([line])
    help_sheet["A1"].font = Font(bold=True, size=14)
    help_sheet.column_dimensions["A"].width = 150

    groups_ws = _sheet(workbook, GROUPS, GROUP_COLUMNS)
    students_ws = _sheet(workbook, STUDENTS, STUDENT_COLUMNS)
    curators_ws = _sheet(workbook, CURATORS, CURATOR_COLUMNS)
    if with_data and db is not None:
        departments = {d.id: d.name for d in db.query(Department).all()}
        groups = db.query(StudyGroup).order_by(StudyGroup.code).all()
        for g in groups:
            append_row(groups_ws, [g.code, g.course, departments.get(g.department_id, ""),
                                   "договор" if g.funding == "contract" else "бюджет", "да" if g.is_active else "нет"])
        code_by_id = {g.id: g.code for g in groups}
        students = (db.query(Student).filter(Student.status != StudentStatus.EXPELLED)
                    .order_by(Student.last_name, Student.first_name).all())
        for s in sorted(students, key=lambda s: (code_by_id.get(s.study_group_id, ""), s.last_name, s.first_name)):
            append_row(students_ws, [s.last_name, s.first_name, s.middle_name or "", code_by_id.get(s.study_group_id, ""),
                                     _STATUS_TEXT[s.status], s.enrolled_at.strftime("%d.%m.%Y")])
        assignments = (db.query(CuratorAssignment).filter(CuratorAssignment.end_date.is_(None)).all())
        for a in sorted(assignments, key=lambda a: (code_by_id.get(a.study_group_id, ""), a.user.full_name)):
            append_row(curators_ws, [a.user.full_name, code_by_id.get(a.study_group_id, ""),
                                     "заместитель" if a.role_type == AssignmentRole.DEPUTY else "куратор",
                                     a.start_date.strftime("%d.%m.%Y")])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
