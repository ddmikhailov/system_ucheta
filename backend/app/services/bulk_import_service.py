"""Разовый импорт посещаемости за сентябрь 2026 из таблиц, которые куратор
вёл до запуска платформы (см. обсуждение в чате — набор из двух Excel-файлов
«1 курс» и «2-4 курс», один лист на группу).

Не для повторного использования как общий механизм — формат этих конкретных
таблиц не описан нигде, кроме этого модуля и переписки. После разового
импорта эндпоинт, который его вызывает, планируется удалить."""
import datetime
import io
from dataclasses import dataclass, field

MARK_CODES = {"о", "и", "б", "з", "п", "р", "у", "н"}

# Разночтения из реальных таблиц, разобранные вручную с администратором
# (см. обсуждение в чате при импорте сентября 2026) — не общий словарь
# алиасов, а фиксированные решения по конкретным встретившимся кодам.
# None означает "не записывать отметку вообще" (не найдено основания).
CODE_ALIASES: dict[str, str | None] = {
    "-": None,  # у куратора не было повода поставить отметку
    "3": "з",  # опечатка/визуальная похожесть с "з" (заявление)
    "на отчисление": None,  # решение принято, но приказа ещё нет — отметку не ставим
}


@dataclass
class ParsedGroupSheet:
    group_code: str
    # ФИО -> {день месяца -> код}
    marks_by_student: dict[str, dict[int, str]] = field(default_factory=dict)
    unrecognized: list[tuple[str, int, str]] = field(default_factory=list)
    error: str | None = None


def _patch_openpyxl_for_nonstandard_border_styles() -> None:
    """Файлы сохранены с border-style 'solid', которого нет в спецификации
    OOXML (openpyxl по умолчанию падает на чтении). Разрешаем его явно —
    сам border нас не интересует, только данные ячеек."""
    import openpyxl.styles.borders as borders_mod

    current = set(borders_mod.Side.style.values)
    if "solid" not in current:
        borders_mod.Side.style.values = frozenset(current | {"solid"})


def parse_attendance_workbook(
    file_bytes: bytes, skip_sheets: frozenset[str] = frozenset({"ИТОГ", "Лист1"})
) -> list[ParsedGroupSheet]:
    _patch_openpyxl_for_nonstandard_border_styles()
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    results: list[ParsedGroupSheet] = []

    for sheet_name in wb.sheetnames:
        if sheet_name in skip_sheets:
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        header_idx = next(
            (i for i, row in enumerate(rows) if row and len(row) > 1 and row[1] == "ФИО"), None
        )
        if header_idx is None:
            results.append(ParsedGroupSheet(group_code=sheet_name, error="не найдена строка заголовка (ФИО)"))
            continue

        sheet = ParsedGroupSheet(group_code=sheet_name)
        for row in rows[header_idx + 1:]:
            if not row or len(row) < 2:
                break
            student_no, fio = row[0], row[1]
            # Конец списка студентов: дальше идут служебные строки
            # ("Начало занятий", суммы по дням) или пустые заготовленные
            # строки без ФИО — по номеру их не отличить, а по типу ФИО можно.
            if not isinstance(student_no, int) or not isinstance(fio, str) or not fio.strip():
                break
            fio = " ".join(fio.split())  # схлопнуть двойные пробелы/переносы

            for day in range(1, 32):
                col = 1 + day  # колонка 2 = день 1, ..., колонка 32 = день 31
                if col >= len(row):
                    break
                val = row[col]
                if val is None:
                    continue
                code = str(val).strip().lower()
                if not code:
                    continue
                if code in MARK_CODES:
                    sheet.marks_by_student.setdefault(fio, {})[day] = code
                elif code in CODE_ALIASES:
                    resolved = CODE_ALIASES[code]
                    if resolved is not None:
                        sheet.marks_by_student.setdefault(fio, {})[day] = resolved
                else:
                    sheet.unrecognized.append((fio, day, str(val)))
        results.append(sheet)

    return results


@dataclass
class ImportReport:
    dry_run: bool
    groups_not_found: list[str] = field(default_factory=list)
    students_not_found: list[dict] = field(default_factory=list)
    unrecognized_codes: list[dict] = field(default_factory=list)
    groups_imported: list[dict] = field(default_factory=list)
    marks_created: int = 0
    submissions_created: int = 0
