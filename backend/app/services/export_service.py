import datetime
import io
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import AttendanceMark, MarkCode, StudyGroup
from app.services import calendar_service, stats_service
from app.services.attendance_service import get_active_students

# Стандартные шрифты reportlab (Helvetica) не содержат кириллицу — без TTF
# с поддержкой Unicode текст в PDF превращается в кракозябры или пробелы.
_CYRILLIC_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "C:\\Windows\\Fonts\\arial.ttf",
]
_PDF_FONT_NAME = "Helvetica"
for _candidate in _CYRILLIC_FONT_CANDIDATES:
    if os.path.exists(_candidate):
        pdfmetrics.registerFont(TTFont("UnicodeBody", _candidate))
        _PDF_FONT_NAME = "UnicodeBody"
        break


def build_summary_workbook(
    db: Session,
    date_from: datetime.date,
    date_to: datetime.date,
    department_id: int | None = None,
) -> bytes:
    """Аналог листа ИТОГ: по строке на группу, числа вместо формул."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ИТОГ"

    headers = [
        "Курс", "Куратор", "Группа", "В списке", "Пришло", "Опоздало",
        "Отсутствует всего", "По уваж. причине", "По неуваж. причине", "Процент присутствия",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    stmt = select(StudyGroup).where(StudyGroup.is_active.is_(True))
    if department_id is not None:
        stmt = stmt.where(StudyGroup.department_id == department_id)
    groups = list(db.execute(stmt.order_by(StudyGroup.course, StudyGroup.code)).scalars().all())

    for group in groups:
        stats = stats_service.compute_period_stats(db, date_from, date_to, study_group_id=group.id)
        curator = next(
            (a.user.full_name for a in group.curator_assignments if a.role_type.value == "curator"), ""
        )
        ws.append(
            [
                group.course, curator, group.code, stats.in_list, stats.present, stats.late,
                stats.absent_total, stats.absent_excused, stats.absent_unexcused, stats.percent,
            ]
        )

    for column_cells in ws.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 40)

    _add_group_sheets(wb, db, groups, date_from, date_to)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _add_group_sheets(wb: Workbook, db: Session, groups: list[StudyGroup], date_from: datetime.date, date_to: datetime.date) -> None:
    for group in groups:
        study_days = calendar_service.study_days_between(
            db, date_from, date_to, study_group_id=group.id, course=group.course
        )
        sheet_name = group.code[:31]
        ws = wb.create_sheet(sheet_name)
        ws.append(["№", "ФИО"] + [d.strftime("%d.%m") for d in study_days])
        for cell in ws[1]:
            cell.font = Font(bold=True)

        students = get_active_students(db, group.id, date_to)
        student_ids = [s.id for s in students]

        marks = db.execute(
            select(AttendanceMark.student_id, AttendanceMark.date, MarkCode.code)
            .join(MarkCode, MarkCode.id == AttendanceMark.mark_code_id)
            .where(AttendanceMark.student_id.in_(student_ids), AttendanceMark.date.in_(study_days))
        ).all()
        marks_map: dict[tuple[int, datetime.date], str] = {(m.student_id, m.date): m.code for m in marks}

        for idx, student in enumerate(students, start=1):
            row = [idx, student.full_name]
            for day in study_days:
                row.append(marks_map.get((student.id, day), ""))
            ws.append(row)


def build_signature_pdf(
    db: Session,
    study_group_id: int,
    date_from: datetime.date,
    date_to: datetime.date,
) -> bytes:
    group = db.get(StudyGroup, study_group_id)
    students = get_active_students(db, study_group_id, date_to)
    stats = stats_service.compute_period_stats(db, date_from, date_to, study_group_id=study_group_id)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=1.5 * cm, rightMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = _PDF_FONT_NAME

    elements = [
        Paragraph(f"Табель посещаемости — группа {group.code}", styles["Title"]),
        Paragraph(f"Период: {date_from.strftime('%d.%m.%Y')} — {date_to.strftime('%d.%m.%Y')}", styles["Normal"]),
        Paragraph(
            f"В списке: {stats.in_list}, присутствие: {stats.percent}%",
            styles["Normal"],
        ),
        Spacer(1, 0.5 * cm),
    ]

    data = [["№", "ФИО", "Дней отсутствия (расчёт за период)"]]
    for idx, student in enumerate(students, start=1):
        student_stats = stats_service.compute_period_stats(db, date_from, date_to, student_id=student.id)
        data.append([idx, student.full_name, student_stats.absent_total])

    table = Table(data, colWidths=[1.5 * cm, 12 * cm, 6 * cm])
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("FONTNAME", (0, 0), (-1, -1), _PDF_FONT_NAME),
            ]
        )
    )
    elements.append(table)
    elements.append(Spacer(1, 1.5 * cm))
    elements.append(Paragraph("Куратор: _____________________ / _____________________ /", styles["Normal"]))

    doc.build(elements)
    return buffer.getvalue()
