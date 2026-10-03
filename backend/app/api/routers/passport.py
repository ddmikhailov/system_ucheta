"""Социальный паспорт группы (futures.md, этап 2): сводка из досье.

Видят: куратор — свои группы, зав. отделением и тьютор — своё отделение, админ,
воспитательный отдел, соц. педагог и психолог — весь колледж. Паспорт
группы (с поимённым списком) читает особые данные досье, поэтому пишет
в журнал просмотров досье запись на каждого студента из категорий; сводка
по группам — только числа, ей достаточно записи в журнале аудита."""
import io

from fastapi import APIRouter, Depends, HTTPException, Response, status
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.api.deps import get_current_user
from app.core.time import today_local
from app.db.session import get_db
from app.models import DossierAccessLog, StudyGroup, User
from app.services import passport_service as svc
from app.services.audit_service import log_action

router = APIRouter(prefix="/passport", tags=["passport"])


class CategoryRead(BaseModel):
    key: str
    title: str
    count: int | None  # None — ключ шифрования недоступен
    names: list[str]


class PassportRead(BaseModel):
    group_id: int
    group_code: str
    course: int
    department_name: str
    students_total: int
    minors: int
    adults: int
    birth_date_missing: int
    budget: int
    contract: int
    funding_missing: int
    no_guardians: int
    dossier_empty: int
    special_available: bool
    categories: list[CategoryRead]


class SummaryRow(BaseModel):
    group_id: int
    group_code: str
    course: int
    department_name: str
    students_total: int
    minors: int
    budget: int
    contract: int
    no_guardians: int
    dossier_empty: int
    counts: dict[str, int | None]


class SummaryRead(BaseModel):
    special_available: bool
    categories: list[dict[str, str]]  # [{key, title}] — порядок столбцов
    rows: list[SummaryRow]
    totals: SummaryRow


def _group_or_403(db: Session, user: User, group_id: int) -> StudyGroup:
    group = db.get(StudyGroup, group_id)
    if group is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Группа не найдена")
    if group_id not in {g.id for g in svc.accessible_groups(db, user)}:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Нет доступа к паспорту этой группы")
    return group


def _passport_read(p: svc.GroupPassport) -> PassportRead:
    return PassportRead(
        group_id=p.group.id, group_code=p.group.code, course=p.group.course,
        department_name=p.group.department.name, students_total=p.students_total, minors=p.minors,
        adults=p.adults, birth_date_missing=p.birth_date_missing, budget=p.budget, contract=p.contract,
        funding_missing=p.funding_missing, no_guardians=p.no_guardians, dossier_empty=p.dossier_empty,
        special_available=p.special_available,
        categories=[
            CategoryRead(key=k, title=title, count=p.counts[k], names=p.names[k]) for k, title, _ in svc.CATEGORIES
        ],
    )


def _log_named_view(db: Session, user: User, passport: svc.GroupPassport) -> None:
    """Поимённый паспорт раскрывает особые данные — как открытие досье: запись в журнал
    просмотров на каждого студента, чьи особые данные оказались в ответе."""
    db.add_all(
        DossierAccessLog(student_id=sid, user_id=user.id, included_special=True)
        for sid in sorted(passport.special_student_ids)
    )


@router.get("/group/{group_id}", response_model=PassportRead)
def group_passport(group_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    group = _group_or_403(db, user, group_id)
    passport = svc.build_passport(db, group, today_local(), with_names=True)
    _log_named_view(db, user, passport)
    log_action(db, user, "passport.view", "study_group", str(group.id))
    db.commit()
    return _passport_read(passport)


def _summary(db: Session, user: User, department_id: int | None) -> tuple[list[svc.GroupPassport], bool]:
    today = today_local()
    passports = [svc.build_passport(db, g, today, with_names=False) for g in svc.accessible_groups(db, user, department_id)]
    return passports, all(p.special_available for p in passports)


def _row(p: svc.GroupPassport) -> SummaryRow:
    return SummaryRow(
        group_id=p.group.id, group_code=p.group.code, course=p.group.course,
        department_name=p.group.department.name, students_total=p.students_total, minors=p.minors,
        budget=p.budget, contract=p.contract, no_guardians=p.no_guardians, dossier_empty=p.dossier_empty,
        counts=p.counts,
    )


def _totals(passports: list[svc.GroupPassport]) -> SummaryRow:
    def total(attr: str) -> int:
        return sum(getattr(p, attr) for p in passports)

    counts: dict[str, int | None] = {}
    for key, _, _ in svc.CATEGORIES:
        values = [p.counts[key] for p in passports]
        counts[key] = None if any(v is None for v in values) else sum(values)
    return SummaryRow(
        group_id=0, group_code="Итого", course=0, department_name="", students_total=total("students_total"),
        minors=total("minors"), budget=total("budget"), contract=total("contract"),
        no_guardians=total("no_guardians"), dossier_empty=total("dossier_empty"), counts=counts,
    )


@router.get("/summary", response_model=SummaryRead)
def summary(
    department_id: int | None = None, user: User = Depends(get_current_user), db: Session = Depends(get_db),
):
    passports, special_ok = _summary(db, user, department_id)
    log_action(db, user, "passport.summary", "department", str(department_id or "all"))
    db.commit()
    return SummaryRead(
        special_available=special_ok,
        categories=[{"key": k, "title": t} for k, t, _ in svc.CATEGORIES],
        rows=[_row(p) for p in passports], totals=_totals(passports),
    )


_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _sheet(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"


@router.get("/export")
def export(
    department_id: int | None = None, group_id: int | None = None,
    user: User = Depends(get_current_user), db: Session = Depends(get_db),
):
    """Excel: сводка по группам (только числа); с group_id — ещё лист с поимённым списком этой группы."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    titles = [t for _, t, _ in svc.CATEGORIES]
    _sheet(ws, ["Группа", "Курс", "Отделение", "Студентов", "Несовершеннолетних", "Бюджет", "Договор",
                "Без представителей", "Досье не заполнено"] + titles)

    if group_id is not None:
        group = _group_or_403(db, user, group_id)
        passports = [svc.build_passport(db, group, today_local(), with_names=True)]
        _log_named_view(db, user, passports[0])
    else:
        passports, _ = _summary(db, user, department_id)

    def cells(row: SummaryRow) -> list:
        return [row.group_code, row.course or "", row.department_name, row.students_total, row.minors, row.budget,
                row.contract, row.no_guardians, row.dossier_empty] + [
            "—" if row.counts[k] is None else row.counts[k] for k, _, _ in svc.CATEGORIES
        ]

    for p in passports:
        ws.append(cells(_row(p)))
    if len(passports) > 1:
        ws.append(cells(_totals(passports)))
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for col in "ABCDEFGHI":
        ws.column_dimensions[col].width = 16

    if group_id is not None:
        names_ws = wb.create_sheet("Списки")
        _sheet(names_ws, ["Категория", "Студент"])
        for key, title, _ in svc.CATEGORIES:
            for name in passports[0].names[key]:
                names_ws.append([title, name])
        names_ws.column_dimensions["A"].width = 28
        names_ws.column_dimensions["B"].width = 40

    log_action(db, user, "passport.export", "study_group" if group_id else "department", str(group_id or department_id or "all"))
    db.commit()
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(), media_type=_XLSX,
        headers={"Content-Disposition": 'attachment; filename="social_passport.xlsx"'},
    )
