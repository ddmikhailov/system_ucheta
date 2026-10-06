"""Разовая актуализация посещаемости «Диджитал» из подготовленного xlsx
(листы «Отметки» и «Дни», см. data/attendance_2026-09_10.xlsx — вне git).

Что делает:
  1. стирает ручные отметки и сдачи дней групп из файла за период файла
     (отметки из длительных отсутствий — source=PERIOD — не трогает);
  2. заново сдаёт каждый день со статусом «заполнено» через attendance_service.submit_day
     от имени администратора («задним числом», чтобы считались свод и группа риска);
     дни «не заполнено» не сдаются вовсе — куратор их не вёл;
  3. студентов и группы, которых нет на сервере, пропускает.

Без --commit — пробный прогон: ничего не записывается, печатается отчёт.

    python -m scripts.import_attendance_xlsx PATH.xlsx            # пробный прогон
    python -m scripts.import_attendance_xlsx PATH.xlsx --commit   # запись
"""
import argparse
import collections
import datetime
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.db import base as db_base
from app.models import (
    AttendanceMark,
    DaySubmission,
    MarkSource,
    Role,
    RoleCode,
    Student,
    StudyGroup,
    User,
)
from app.services import attendance_service

FILLED = "заполнено"


def _to_date(value) -> datetime.date:
    return value.date() if isinstance(value, datetime.datetime) else value


def read_workbook(path: str):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    marks = collections.defaultdict(dict)  # (группа, дата) -> {ФИО: код}
    for group, fio, day, code in list(wb["Отметки"].iter_rows(min_row=2, values_only=True)):
        marks[(group, _to_date(day))][fio] = code
    filled = collections.defaultdict(list)  # группа -> [даты «заполнено»]
    all_days = []
    for group, day, status in list(wb["Дни"].iter_rows(min_row=2, values_only=True)):
        all_days.append(_to_date(day))
        if status == FILLED:
            filled[group].append(_to_date(day))
    return marks, filled, min(all_days), max(all_days)


def load(db: Session, path: str, admin: User, commit: bool) -> dict:
    marks, filled, date_from, date_to = read_workbook(path)
    report = {
        "groups_not_found": [], "students_not_found": [], "marks_skipped_inactive": 0,
        "days_submitted": 0, "days_not_submitted": 0, "marks_written": 0,
        "deleted_marks": 0, "deleted_submissions": 0, "kept_period_marks": 0,
    }

    groups = {}
    for code in sorted(filled):
        group = db.query(StudyGroup).filter(StudyGroup.code == code).one_or_none()
        if group is None:
            report["groups_not_found"].append(code)
        else:
            groups[code] = group

    group_ids = [g.id for g in groups.values()]
    student_ids = [
        sid for (sid,) in db.execute(select(Student.id).where(Student.study_group_id.in_(group_ids)))
    ]
    in_range = (AttendanceMark.date >= date_from, AttendanceMark.date <= date_to)
    old_marks = db.execute(
        select(AttendanceMark.source).where(AttendanceMark.student_id.in_(student_ids), *in_range)
    ).scalars().all()
    report["deleted_marks"] = sum(1 for s in old_marks if s == MarkSource.MANUAL)
    report["kept_period_marks"] = sum(1 for s in old_marks if s == MarkSource.PERIOD)
    report["deleted_submissions"] = len(db.execute(
        select(DaySubmission.id).where(
            DaySubmission.study_group_id.in_(group_ids),
            DaySubmission.date >= date_from, DaySubmission.date <= date_to,
        )
    ).all())

    if commit:
        db.execute(delete(AttendanceMark).where(
            AttendanceMark.student_id.in_(student_ids), AttendanceMark.source == MarkSource.MANUAL, *in_range,
        ))
        db.execute(delete(DaySubmission).where(
            DaySubmission.study_group_id.in_(group_ids),
            DaySubmission.date >= date_from, DaySubmission.date <= date_to,
        ))
        db.commit()

    missing = set()
    for code, group in groups.items():
        by_name = {s.full_name: s for s in db.query(Student).filter(Student.study_group_id == group.id)}
        for day in sorted(filled[code]):
            exceptions = []
            active = {s.id for s in attendance_service.get_active_students(db, group.id, day)}
            for fio, mark_code in marks.get((code, day), {}).items():
                student = by_name.get(fio)
                if student is None:
                    missing.add((code, fio))
                elif student.id not in active:
                    report["marks_skipped_inactive"] += 1
                else:
                    exceptions.append(
                        {"student_id": student.id, "mark_code": mark_code, "comment": None, "basis_reference": None}
                    )
            try:
                if commit:
                    attendance_service.submit_day(db, group.id, day, exceptions, admin, today=day)
            except attendance_service.BackdateNotAllowed:
                report["days_not_submitted"] += 1  # по календарю сервера день нерабочий
                continue
            report["days_submitted"] += 1
            report["marks_written"] += len(exceptions)
    report["students_not_found"] = sorted(missing)
    if not commit:
        db.rollback()
    return report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("path")
    parser.add_argument("--commit", action="store_true")
    args = parser.parse_args()
    db = db_base.SessionLocal()
    try:
        admin = db.query(User).join(Role).filter(Role.code == RoleCode.ADMIN.value).order_by(User.id).first()
        if admin is None:
            raise RuntimeError("Администратор не найден")
        report = load(db, args.path, admin, args.commit)
    finally:
        db.close()
    print("ЗАПИСЬ" if args.commit else "ПРОБНЫЙ ПРОГОН (ничего не записано)")
    for key, value in report.items():
        if isinstance(value, list):
            print(f"{key}: {len(value)}")
            for item in value:
                print("   ", item)
        else:
            print(f"{key}: {value}")


if __name__ == "__main__":
    main()
