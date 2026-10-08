"""Единая загрузка контингента из Excel-шаблона: синтетические ФИО."""
import datetime
import io

import openpyxl
import pytest

import scripts.import_contingent as cli
from app.core.security import verify_password
from app.core.xlsx import append_row
from app.models import (
    AssignmentRole, AuditLog, CuratorAssignment, Department, Student, StudentGroupMembership, StudentStatus, StudyGroup, User,
)
from app.services import contingent_import_service as svc

TODAY = datetime.date(2026, 10, 1)
GROUPS_HEAD = ["Код группы", "Курс", "Отделение", "Финансирование (бюджет/договор)", "Группа действует (да/нет)"]
STUDENTS_HEAD = ["Фамилия", "Имя", "Отчество", "Группа", "Статус (обучается / академ. отпуск / отчислен)", "Дата зачисления"]
CURATORS_HEAD = ["ФИО куратора", "Группа", "Роль (куратор / заместитель)", "С какой даты"]


def book(groups=None, students=None, curators=None, extra_sheet=False) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, head, rows in (("Группы", GROUPS_HEAD, groups), ("Студенты", STUDENTS_HEAD, students), ("Кураторы", CURATORS_HEAD, curators)):
        if rows is None:
            continue
        ws = wb.create_sheet(title)
        ws.append(head)
        for row in rows:
            append_row(ws, row)  # как делает платформа: «=…» остаётся текстом
    if extra_sheet:
        wb.create_sheet("Инструкция").append(["игнорируется"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def options(**kw) -> svc.Options:
    return svc.Options(today=TODAY, **kw)


def load(db, content: bytes, apply: bool = True, **kw) -> svc.Report:
    report = svc.process(db, svc.parse_workbook(content), options(**kw))
    if apply and not report.blocked:
        db.commit()
    else:
        db.rollback()
    return report


@pytest.fixture()
def digital(seeded, db):
    return db.query(Department).filter_by(name="Диджитал").one()


FULL = dict(
    groups=[["ТЕСТ111", 1, "Диджитал", "бюджет", "да"], ["ТЕСТ212", 2, "Диджитал", "договор", "да"]],
    students=[["Тестов", "Иван", "Петрович", "ТЕСТ111", "обучается", "01.09.2026"],
              ["Образцова", "Анна", "", "ТЕСТ111", "", ""],
              ["Пробный", "Олег", "Игоревич", "ТЕСТ212", "академ. отпуск", ""]],
    curators=[["Куратова Мария Ивановна", "ТЕСТ111", "куратор", ""], ["Куратова Мария Ивановна", "ТЕСТ212", "", "05.09.2026"]],
)


def test_creates_groups_students_and_curators(digital, db):
    report = load(db, book(**FULL, extra_sheet=True))

    assert not report.errors and report.counts == {
        "групп создано": 2, "студентов создано": 3, "кураторов создано": 1, "назначений создано": 2,
    }
    group = db.query(StudyGroup).filter_by(code="ТЕСТ212").one()
    assert group.course == 2 and group.funding == "contract" and group.department_id == digital.id
    leave = db.query(Student).filter_by(last_name="Пробный").one()
    assert leave.status == StudentStatus.ACADEMIC_LEAVE and leave.enrolled_at == TODAY
    assert db.query(Student).filter_by(last_name="Тестов").one().enrolled_at == datetime.date(2026, 9, 1)
    assert db.query(StudentGroupMembership).count() == 3
    curator = db.query(User).filter_by(full_name="Куратова Мария Ивановна").one()
    assert curator.username == "kuratova.m" and curator.password_hash is None and curator.department_id == digital.id
    starts = {a.study_group.code: a.start_date for a in db.query(CuratorAssignment).filter_by(user_id=curator.id)}
    assert starts == {"ТЕСТ111": TODAY, "ТЕСТ212": datetime.date(2026, 9, 5)}


def test_second_load_and_preview_change_nothing(digital, db):
    content = book(**FULL)
    preview = load(db, content, apply=False)
    assert preview.counts["студентов создано"] == 3 and db.query(Student).count() == 0  # предпросмотр откатывается

    load(db, content)
    again = load(db, content)
    assert again.counts == {} and not again.changes and not again.errors
    assert db.query(Student).count() == 3 and db.query(CuratorAssignment).count() == 2


def test_any_error_blocks_everything_and_is_listed_with_row_numbers(digital, db):
    content = book(
        groups=[["ТЕСТ111", 1, "Диджитал", "", ""], ["НОВАЯ", "", "", "", ""], ["КРИВАЯ", 9, "Диджитал", "", ""],
                ["ФИН", 1, "Диджитал", "частично", ""], ["ЧУЖАЯ", 1, "Несуществующее", "", ""]],
        students=[["Иванов", "Иван", "", "НЕТТАКОЙ", "", ""], ["", "Имя", "", "ТЕСТ111", "", ""],
                  ["Статусов", "Сергей", "", "ТЕСТ111", "умер", ""], ["Датов", "Дмитрий", "", "ТЕСТ111", "", "вчера"]],
        curators=[["Один", "ТЕСТ111", "", ""], ["Куратор Ролевой", "ТЕСТ111", "директор", ""]],
    )
    report = load(db, content)
    messages = {(e.sheet, e.row): e.message for e in report.errors}
    assert "нужны курс и отделение" in messages[("Группы", 3)]
    assert "от 1 до 6" in messages[("Группы", 4)]
    assert "«бюджет» или «договор»" in messages[("Группы", 5)]
    assert "не найдено на платформе" in messages[("Группы", 6)]
    assert "не найдена" in messages[("Студенты", 2)] or "НЕТТАКОЙ" in messages[("Студенты", 2)]
    assert "нужны фамилия, имя и группа" in messages[("Студенты", 3)]
    assert "статус «умер»" in messages[("Студенты", 4)] and "нужна дата" in messages[("Студенты", 5)]
    assert "нужны ФИО куратора" in messages[("Кураторы", 2)] and "роль «директор»" in messages[("Кураторы", 3)]
    assert db.query(StudyGroup).count() == 0 and db.query(Student).count() == 0  # ничего не записано


def test_duplicates_in_file_are_errors(digital, db):
    content = book(groups=[["ТЕСТ111", 1, "Диджитал", "", ""], ["тест111", 1, "Диджитал", "", ""]],
                   students=[["Двойной", "Иван", "", "ТЕСТ111", "", ""], ["двойной", "иван", "", "ТЕСТ111", "", ""]])
    errors = load(db, content).errors
    assert any("дважды" in e.message for e in errors if e.sheet == "Группы")
    assert any("дважды" in e.message for e in errors if e.sheet == "Студенты")


def test_unknown_department_error_or_creation(digital, db):
    content = book(groups=[["ТЕСТ111", 1, "Новое отделение", "", ""]])
    assert load(db, content).errors
    report = load(db, content, create_departments=True)
    assert not report.errors and report.counts["отделений создано"] == 1
    assert db.query(StudyGroup).filter_by(code="ТЕСТ111").one().department.name == "Новое отделение"


def test_transfer_status_change_and_return_of_expelled(digital, db):
    load(db, book(groups=FULL["groups"], students=[
        ["Перевёдов", "Пётр", "", "ТЕСТ111", "", ""], ["Отчисленов", "Олег", "", "ТЕСТ111", "отчислен", ""],
        ["Академов", "Антон", "", "ТЕСТ111", "", ""]]))
    assert db.query(Student).filter_by(last_name="Отчисленов").one().status == StudentStatus.EXPELLED

    report = load(db, book(students=[
        ["перевёдов", "пётр", "", "ТЕСТ212", "", ""],            # другая группа, регистр/ё не важны → перевод
        ["Отчисленов", "Олег", "", "ТЕСТ111", "", ""],           # указан без статуса → возвращён
        ["Академов", "Антон", "", "ТЕСТ111", "академ. отпуск", ""]]))
    assert report.counts == {"студентов переведено": 1, "студентов возвращено из отчисленных": 1, "статус студента изменён": 1}
    db.expire_all()
    moved = db.query(Student).filter_by(last_name="Перевёдов").one()
    assert moved.study_group.code == "ТЕСТ212"
    memberships = db.query(StudentGroupMembership).filter_by(student_id=moved.id).order_by(StudentGroupMembership.id).all()
    assert [m.study_group.code for m in memberships] == ["ТЕСТ111", "ТЕСТ212"]
    assert memberships[0].end_date == TODAY - datetime.timedelta(days=1)
    assert db.query(Student).filter_by(last_name="Отчисленов").one().status == StudentStatus.STUDYING
    assert db.query(Student).filter_by(last_name="Академов").one().status == StudentStatus.ACADEMIC_LEAVE
    assert db.query(Student).count() == 3  # дублей нет


def test_students_missing_from_file_are_kept_by_default_and_only_listed_groups_can_expel(digital, db):
    load(db, book(groups=FULL["groups"], students=[
        ["Остающийся", "Иван", "", "ТЕСТ111", "", ""], ["Ушедший", "Пётр", "", "ТЕСТ111", "", ""],
        ["Другогрупповой", "Олег", "", "ТЕСТ212", "", ""]]))
    one_student = book(students=[["Остающийся", "Иван", "", "ТЕСТ111", "", ""]])

    keep = load(db, one_student)
    assert not keep.counts and any("нет 1 студентов" in w for w in keep.warnings)

    expel = load(db, one_student, absent_students="expel")
    assert expel.counts == {"студентов отмечено отчисленными": 1}
    db.expire_all()
    assert db.query(Student).filter_by(last_name="Ушедший").one().status == StudentStatus.EXPELLED
    assert db.query(Student).filter_by(last_name="Ушедший").one().left_at == TODAY
    # группа ТЕСТ212 в файле не упоминалась — её студент не тронут
    assert db.query(Student).filter_by(last_name="Другогрупповой").one().status == StudentStatus.STUDYING


def test_mass_expulsion_needs_confirmation(digital, db):
    load(db, book(groups=FULL["groups"], students=[[f"Студент{i:02d}", "Тест", "", "ТЕСТ111", "", ""] for i in range(20)]))
    partial = book(students=[[f"Студент{i:02d}", "Тест", "", "ТЕСТ111", "", ""] for i in range(5)])

    report = load(db, partial, absent_students="expel")
    assert report.needs_confirmation and "15 из 20" in report.needs_confirmation
    db.expire_all()
    assert db.query(Student).filter_by(status=StudentStatus.EXPELLED).count() == 0  # заблокировано

    report = load(db, partial, absent_students="expel", confirm_large=True)
    assert not report.needs_confirmation and report.counts["студентов отмечено отчисленными"] == 15


def test_curator_conflicts_replace_and_deputy(digital, db):
    load(db, book(groups=FULL["groups"], curators=[["Первый Куратор", "ТЕСТ111", "", "01.09.2026"]]))

    keep = load(db, book(curators=[["Второй Куратор", "ТЕСТ111", "", "10.10.2026"]]))
    assert keep.counts["назначений пропущено (уже есть другой)"] == 1
    assert any("уже назначен Первый Куратор" in w for w in keep.warnings)
    assert db.query(CuratorAssignment).filter(CuratorAssignment.end_date.is_(None)).count() == 1

    replaced = load(db, book(curators=[["Второй Куратор", "ТЕСТ111", "", "10.10.2026"]]), replace_curators=True)
    assert replaced.counts["назначений закрыто (замена)"] == 1
    old = db.query(CuratorAssignment).filter_by(end_date=datetime.date(2026, 10, 9)).one()
    assert old.user.full_name == "Первый Куратор"
    assert db.query(CuratorAssignment).filter(CuratorAssignment.end_date.is_(None)).one().user.full_name == "Второй Куратор"

    load(db, book(curators=[["Заместитель Зам", "ТЕСТ111", "заместитель", ""]]))
    deputy = db.query(CuratorAssignment).filter_by(role_type=AssignmentRole.DEPUTY).one()
    assert deputy.user.role.code == "deputy_curator" and deputy.end_date is None


def test_temporary_passwords_only_when_asked_and_only_for_curators_without_one(digital, db):
    content = book(groups=FULL["groups"], curators=FULL["curators"])
    load(db, content)
    assert db.query(User).filter_by(username="kuratova.m").one().password_hash is None

    report = load(db, content, issue_passwords=True)
    assert len(report.credentials) == 1 and report.counts == {"временных паролей выдано": 1}
    credential = report.credentials[0]
    user = db.query(User).filter_by(username=credential.username).one()
    assert user.must_change_password and verify_password(credential.password, user.password_hash)
    assert load(db, content, issue_passwords=True).credentials == []  # второй раз — не меняем и не выдаём


def test_template_with_data_round_trips_without_changes(digital, db):
    load(db, book(**FULL))
    template = svc.build_template(db, with_data=True)
    wb = openpyxl.load_workbook(io.BytesIO(template))
    assert wb.sheetnames == ["Инструкция", "Группы", "Студенты", "Кураторы"]
    assert wb["Студенты"].max_row == 4 and wb["Кураторы"].max_row == 3
    again = load(db, template)
    assert again.counts == {} and not again.errors and not again.warnings

    empty = openpyxl.load_workbook(io.BytesIO(svc.build_template(db, with_data=False)))
    assert empty["Студенты"].max_row == 1


def test_garbage_and_wrong_files_are_rejected_clearly(digital, db):
    with pytest.raises(svc.ImportFileError, match="не xlsx-файл"):
        svc.parse_workbook(b"not an excel file")
    wb = openpyxl.Workbook()
    wb.active.title = "Чужой лист"
    out = io.BytesIO()
    wb.save(out)
    with pytest.raises(svc.ImportFileError, match="нет ни одного из листов"):
        svc.parse_workbook(out.getvalue())
    wb = openpyxl.Workbook()
    wb.active.title = "Студенты"
    wb.active.append(["Фамилия", "Имя"])
    out = io.BytesIO()
    wb.save(out)
    with pytest.raises(svc.ImportFileError, match="обязательные колонки"):
        svc.parse_workbook(out.getvalue())


def test_formula_like_names_stay_text_in_template(digital, db):
    load(db, book(groups=FULL["groups"][:1], students=[["=1+1", "Хитрый", "", "ТЕСТ111", "", ""]]))
    ws = openpyxl.load_workbook(io.BytesIO(svc.build_template(db)))["Студенты"]
    assert ws["A2"].data_type != "f" and str(ws["A2"].value).lstrip("'") == "=1+1"


# ---- HTTP ----

def _post(client, headers, path, content, **params):
    query = "&".join(f"{k}={str(v).lower() if isinstance(v, bool) else v}" for k, v in params.items())
    return client.post(f"/contingent-import/{path}" + (f"?{query}" if query else ""), headers=headers, content=content)


def test_http_preview_does_not_write_apply_writes_and_audit_has_no_passwords(client, admin_headers, digital, db):
    content = book(**FULL)
    preview = _post(client, admin_headers, "preview", content)
    assert preview.status_code == 200
    body = preview.json()
    assert body["counts"]["студентов создано"] == 3 and body["sheets"] == ["Группы", "Студенты", "Кураторы"]
    assert body["errors"] == [] and body["total_changes"] >= 6
    assert db.query(Student).count() == 0

    applied = _post(client, admin_headers, "apply", content, issue_passwords=True)
    assert applied.status_code == 200, applied.text
    data = applied.json()
    assert len(data["credentials"]) == 1 and data["counts"]["временных паролей выдано"] == 1
    db.expire_all()
    assert db.query(Student).count() == 3
    entry = db.query(AuditLog).filter_by(action="contingent.import").one()
    assert data["credentials"][0]["password"] not in (entry.new_value or "") and "студентов создано: 3" in entry.new_value


def test_http_apply_refuses_files_with_errors_or_unconfirmed_mass_expulsion(client, admin_headers, digital, db):
    bad = _post(client, admin_headers, "apply", book(students=[["Иванов", "Иван", "", "НЕТТАКОЙ", "", ""]]))
    assert bad.status_code == 400 and "запись не выполнена" in bad.json()["detail"]

    load(db, book(groups=FULL["groups"], students=[[f"Студент{i:02d}", "Тест", "", "ТЕСТ111", "", ""] for i in range(20)]))
    partial = book(students=[[f"Студент{i:02d}", "Тест", "", "ТЕСТ111", "", ""] for i in range(3)])
    refused = _post(client, admin_headers, "apply", partial, absent_students="expel")
    assert refused.status_code == 409 and "17 из 20" in refused.json()["detail"]
    ok = _post(client, admin_headers, "apply", partial, absent_students="expel", confirm_large=True)
    assert ok.status_code == 200


def test_http_file_checks(client, admin_headers, digital):
    assert _post(client, admin_headers, "preview", b"").status_code == 400
    assert _post(client, admin_headers, "preview", b"xxx").status_code == 400
    assert _post(client, admin_headers, "preview", b"x" * (5 * 1024 * 1024 + 1)).status_code == 413
    assert client.post("/contingent-import/preview?absent_students=delete", headers=admin_headers, content=b"x").status_code == 422


def test_http_template_download(client, admin_headers, digital, db):
    r = client.get("/contingent-import/template", headers=admin_headers)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames[1:] == ["Группы", "Студенты", "Кураторы"]
    empty = client.get("/contingent-import/template?with_data=false", headers=admin_headers)
    assert "empty" in empty.headers["content-disposition"]


@pytest.mark.parametrize("headers_fixture", ["curator_headers", "dept_head_headers", "edu_department_headers", "meal_manager_headers"])
def test_http_only_administrator_may_use_it(client, request, headers_fixture, digital):
    headers = request.getfixturevalue(headers_fixture)
    for method, path in (("get", "/contingent-import/template"), ("post", "/contingent-import/preview"), ("post", "/contingent-import/apply")):
        assert getattr(client, method)(path, headers=headers).status_code == 403, (headers_fixture, path)
    assert client.get("/contingent-import/template").status_code == 401


# ---- консольная команда ----

def test_cli_preview_apply_report_and_passwords(digital, db, tmp_path, monkeypatch):
    path = tmp_path / "in.xlsx"
    path.write_bytes(book(**FULL))
    monkeypatch.setattr(cli.db_base, "SessionLocal", lambda: _TestSession(db))

    dry = cli.run(str(path), False, options())
    assert dry.counts["студентов создано"] == 3 and db.query(Student).count() == 0

    report_file, passwords = tmp_path / "r.txt", tmp_path / "p.csv"
    done = cli.run(str(path), True, options(issue_passwords=True), issue_passwords_file=str(passwords), report_file=str(report_file))
    assert not done.blocked and db.query(Student).count() == 3
    assert "ЗАПИСЬ" in report_file.read_text(encoding="utf-8") and "kuratova.m" in passwords.read_text(encoding="utf-8-sig")
    assert oct(passwords.stat().st_mode & 0o777) == "0o600"

    blocked_file = tmp_path / "bad.xlsx"
    blocked_file.write_bytes(book(students=[["Иванов", "Иван", "", "НЕТТАКОЙ", "", ""]]))
    blocked = cli.run(str(blocked_file), True, options())
    assert blocked.blocked and db.query(Student).count() == 3


class _TestSession:
    """Обёртка: скрипт закрывает сессию в конце, а тесту она нужна дальше."""

    def __init__(self, session):
        self._s = session

    def __getattr__(self, name):
        return getattr(self._s, name)

    def close(self):
        pass
