"""Загрузка досье из Excel: шаблон, предпросмотр ошибок, применение."""
import datetime
import io

from openpyxl import Workbook, load_workbook

from app.core import field_crypto
from app.core.config import get_settings
from app.models import Department, Student, StudentGuardian, StudentProfile, StudyGroup
from app.services import dossier_import_service as svc

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _file(rows: list[dict]) -> bytes:
    """Файл по шаблону: значения по заголовкам, остальное пусто."""
    headers = svc._headers()
    wb = Workbook()
    ws = wb.active
    ws.title = "Досье"
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client, path, headers, content):
    return client.post(path, headers={**headers, "Content-Type": XLSX}, content=content)


def _who(student: Student) -> dict:
    return {"Группа": student.study_group.code, "Фамилия": student.last_name, "Имя": student.first_name,
            "Отчество": student.middle_name}


def test_template_has_headers_and_only_names_of_accessible_students(client, curator_headers, admin_headers, imported, db):
    assert client.get("/dossier-import/template", headers=curator_headers).status_code == 403  # куратор — через карточку
    r = client.get("/dossier-import/template", headers=admin_headers)
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0][: len(svc._headers())]) == svc._headers()
    assert len(rows) - 1 == db.query(Student).count()
    assert all(not any(row[4:]) for row in rows[1:])  # кроме имён — пусто, данных студентов нет


def test_preview_and_apply_updates_profile_guardians_and_special(client, admin_headers, imported, db):
    s = db.query(Student).first()
    content = _file([{
        **_who(s), "Дата рождения": "05.03.2008", "Финансирование (бюджет/договор)": "Бюджет",
        "Телефон": "+7 900 111-22-33", "Представитель 1: ФИО": "Иванова А.", "Представитель 1: Кем приходится": "мать",
        "Сирота (да/нет)": "да", "Здоровье": "астма", "ОВЗ (да/нет)": "нет",
    }])
    preview = _post(client, "/dossier-import/preview", admin_headers, content).json()
    assert (preview["total"], preview["ready"], preview["with_errors"]) == (1, 1, 0)
    assert db.get(StudentProfile, s.id) is None  # предпросмотр ничего не пишет

    r = _post(client, "/dossier-import/apply", admin_headers, content)
    assert r.status_code == 200 and r.json() == {"updated": 1, "skipped_with_errors": 0}
    db.expire_all()
    profile = db.get(StudentProfile, s.id)
    assert profile.birth_date == datetime.date(2008, 3, 5) and profile.funding == "budget"
    assert profile.phone == "+7 900 111-22-33"
    special = field_crypto.decrypt_json(profile.special_enc)
    assert special["is_orphan"] is True and special["health_note"] == "астма" and special["has_ovz"] is False
    guardians = db.query(StudentGuardian).filter(StudentGuardian.student_id == s.id).all()
    assert [(g.full_name, g.relation, g.is_primary) for g in guardians] == [("Иванова А.", "мать", True)]


def test_empty_cells_do_not_erase_and_reimport_is_idempotent(client, admin_headers, imported, db):
    s = db.query(Student).first()
    first = _file([{**_who(s), "Телефон": "123", "Сирота (да/нет)": "да"}])
    _post(client, "/dossier-import/apply", admin_headers, first)
    second = _file([{**_who(s), "E-mail": "a@b.ru"}])  # телефон и особые не указаны
    _post(client, "/dossier-import/apply", admin_headers, second)
    db.expire_all()
    profile = db.get(StudentProfile, s.id)
    assert profile.phone == "123" and profile.email == "a@b.ru"
    assert field_crypto.decrypt_json(profile.special_enc)["is_orphan"] is True
    again = _post(client, "/dossier-import/preview", admin_headers, second).json()
    assert again["ready"] == 1  # «готово» = в строке есть данные; повторная загрузка безопасна
    assert again["with_errors"] == 0


def test_errors_are_reported_and_bad_rows_skipped(client, admin_headers, imported, db):
    a, b = db.query(Student).limit(2).all()
    content = _file([
        {**_who(a), "Телефон": "111"},
        {**_who(b), "Дата рождения": "не дата", "Финансирование (бюджет/договор)": "стипендия", "Сирота (да/нет)": "может"},
        {"Группа": "НЕТТАКОЙ", "Фамилия": "Х", "Имя": "Х"},
        {**_who(a), "Фамилия": "Несуществующий", "Телефон": "1"},
        {**_who(a), "Представитель 1: Кем приходится": "мать"},
    ])
    preview = _post(client, "/dossier-import/preview", admin_headers, content).json()
    assert (preview["total"], preview["ready"], preview["with_errors"]) == (5, 1, 4)
    errors = {r["row"]: r["errors"] for r in preview["rows"] if r["errors"]}
    assert len(errors[3]) == 3  # дата, финансирование, да/нет
    assert errors[4] == ["группа не найдена"] and errors[5] == ["студент не найден в этой группе"]
    assert "не указано ФИО" in errors[6][0]
    assert preview["rows"][0]["errors"]  # ошибки показываются первыми

    r = _post(client, "/dossier-import/apply", admin_headers, content).json()
    assert r == {"updated": 1, "skipped_with_errors": 4}
    db.expire_all()
    assert db.get(StudentProfile, a.id).phone == "111" and db.get(StudentProfile, b.id) is None


def test_special_cells_rejected_without_encryption_key(client, admin_headers, imported, db, monkeypatch):
    monkeypatch.setattr(get_settings(), "dossier_encryption_key", "")
    s = db.query(Student).first()
    preview = _post(client, "/dossier-import/preview", admin_headers,
                    _file([{**_who(s), "Телефон": "1", "Здоровье": "х"}])).json()
    assert preview["with_errors"] == 1 and "ключ шифрования" in preview["rows"][0]["errors"][0]
    ok = _post(client, "/dossier-import/preview", admin_headers, _file([{**_who(s), "Телефон": "1"}])).json()
    assert ok["ready"] == 1


def test_dept_head_import_limited_to_own_department(client, dept_head_headers, dept_head_user, imported, db):
    other = Department(name="Чужое")
    db.add(other)
    db.flush()
    foreign_group = StudyGroup(code="FRG-1", course=1, department_id=other.id)
    db.add(foreign_group)
    db.flush()
    foreign = Student(last_name="Чужой", first_name="Студент", study_group_id=foreign_group.id,
                      enrolled_at=datetime.date(2026, 9, 1))
    db.add(foreign)
    db.commit()
    own = db.query(Student).filter(Student.study_group_id != foreign_group.id).first()

    template = client.get("/dossier-import/template", headers=dept_head_headers)
    names = {row[1] for row in load_workbook(io.BytesIO(template.content)).worksheets[0].iter_rows(min_row=2, values_only=True)}
    assert "Чужой" not in names

    content = _file([{**_who(own), "Телефон": "1"}, {**_who(foreign), "Телефон": "2"}])
    preview = _post(client, "/dossier-import/preview", dept_head_headers, content).json()
    assert preview["ready"] == 1 and preview["with_errors"] == 1
    assert preview["rows"][0]["errors"] == ["группа не из вашего отделения"]


def test_bad_files_are_rejected(client, admin_headers):
    assert _post(client, "/dossier-import/preview", admin_headers, b"not an excel").status_code == 400
    wrong = Workbook()
    wrong.active.append(["a", "b"])
    buf = io.BytesIO()
    wrong.save(buf)
    r = _post(client, "/dossier-import/preview", admin_headers, buf.getvalue())
    assert r.status_code == 400 and "шаблон" in r.json()["detail"]
    assert _post(client, "/dossier-import/preview", admin_headers, b"").status_code == 400


def _zip(entries: dict[str, bytes]) -> bytes:
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return buf.getvalue()


def test_zip_bomb_is_rejected_before_parsing(client, admin_headers):
    """Файл ≤ 5 МБ, но после распаковки — 200 МБ нулей: не должен дойти до openpyxl."""
    bomb = _zip({"xl/worksheets/sheet1.xml": b"\0" * (200 * 1024 * 1024)})
    assert len(bomb) < 5 * 1024 * 1024
    r = _post(client, "/dossier-import/preview", admin_headers, bomb)
    assert r.status_code == 400 and "после распаковки" in r.json()["detail"]
    assert _post(client, "/dossier-import/apply", admin_headers, bomb).status_code == 400


def test_archive_with_too_many_entries_is_rejected(client, admin_headers):
    many = _zip({f"f{i}.xml": b"x" for i in range(300)})
    r = _post(client, "/dossier-import/preview", admin_headers, many)
    assert r.status_code == 400 and "внутренних файлов" in r.json()["detail"]


def test_real_size_is_counted_even_if_the_header_lies(monkeypatch):
    """В заголовке zip размер можно занизить — поэтому считаются реально распакованные байты."""
    import struct

    from app.core import xlsx

    bomb = bytearray(_zip({"a.xml": b"\0" * (3 * 1024 * 1024)}))
    # file_size в центральном каталоге (смещение 24 от сигнатуры PK\x01\x02) подменяем на 10 байт
    pos = bomb.rindex(b"PK\x01\x02")
    bomb[pos + 24:pos + 28] = struct.pack("<I", 10)
    monkeypatch.setattr(xlsx, "MAX_UNCOMPRESSED_BYTES", 1024 * 1024)
    try:
        xlsx.check_zip_safety(bytes(bomb))
    except xlsx.UnsafeArchiveError as exc:
        assert "после распаковки" in str(exc) or "повреждён" in str(exc)
    else:
        raise AssertionError("архив с подделанным размером прошёл проверку")


def test_normal_template_passes_the_zip_check():
    from app.core import xlsx

    xlsx.check_zip_safety(_file([{"Группа": "X", "Фамилия": "Я", "Имя": "Я"}]))


def test_too_many_rows_and_sparse_huge_sheet_are_rejected(client, admin_headers):
    headers = svc._headers()
    wb = Workbook()
    ws = wb.active
    ws.title = "Досье"
    ws.append(headers)
    for _ in range(svc.MAX_SCANNED_ROWS + 5):
        ws.append(["x"])
    buf = io.BytesIO()
    wb.save(buf)
    r = _post(client, "/dossier-import/preview", admin_headers, buf.getvalue())
    assert r.status_code == 400 and "Слишком много строк" in r.json()["detail"]

    # лист, объявивший огромную область одной далёкой ячейкой, не обходится целиком
    wb = Workbook()
    ws = wb.active
    ws.title = "Досье"
    ws.append(headers)
    ws.cell(row=1_000_000, column=1, value="далеко")
    buf = io.BytesIO()
    wb.save(buf)
    r = _post(client, "/dossier-import/preview", admin_headers, buf.getvalue())
    assert r.status_code == 400 and "Слишком много строк" in r.json()["detail"]
