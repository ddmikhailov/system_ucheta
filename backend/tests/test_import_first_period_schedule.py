import datetime

import openpyxl

from app.models import DaySubmission
from app.services import attendance_service
from scripts import import_first_period_schedule as importer

D1 = datetime.date(2026, 9, 21)  # понедельник
D2 = datetime.date(2026, 9, 22)


def _write_schedule(path, group_code, lessons_by_day):
    """Синтетическое расписание той же формы: блок дня — 5 колонок, на пару — две строки уроков."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = importer.SHEET
    days = sorted(lessons_by_day)
    header = [None] * (5 + 5 * len(days))
    for i, day in enumerate(days):
        header[5 + 5 * i] = f"{day:%d.%m.%Y}\nДЕНЬ"
    for _ in range(3):
        ws.append([])
    ws.append(header)
    for _ in range(4):
        ws.append([])
    for lesson in range(1, 13):
        row = [None] * len(header)
        row[2] = str(lesson)
        if lesson == 1:
            row[0] = group_code
        for i, day in enumerate(days):
            if (lesson + 1) // 2 in lessons_by_day[day]:
                row[5 + 5 * i] = "Дисциплина" if lesson % 2 else "Преподаватель"
        ws.append(row)
    wb.save(path)


def test_fills_first_pair_only_for_empty_by_default(db, curator_group, curator_user, tmp_path):
    for day in (D1, D2):
        attendance_service.submit_day(db, curator_group.id, day, [], curator_user, today=day)
    sub2 = db.query(DaySubmission).filter_by(study_group_id=curator_group.id, date=D2).one()
    sub2.first_period = 4
    db.commit()
    path = tmp_path / "s.xlsx"
    _write_schedule(path, curator_group.code, {D1: {2, 3}, D2: {1}})

    report = importer.load(db, [str(path)], commit=True)

    subs = {s.date: s.first_period for s in db.query(DaySubmission).filter_by(study_group_id=curator_group.id)}
    assert subs == {D1: 2, D2: 4}
    assert report["updated"] == 1 and report["kept_existing"] == 1

    importer.load(db, [str(path)], commit=True, overwrite=True)
    db.expire_all()
    assert db.query(DaySubmission).filter_by(study_group_id=curator_group.id, date=D2).one().first_period == 1


def test_dry_run_and_unknown_group_write_nothing(db, curator_group, curator_user, tmp_path):
    attendance_service.submit_day(db, curator_group.id, D1, [], curator_user, today=D1)
    path = tmp_path / "s.xlsx"
    _write_schedule(path, "НЕТ-00", {D1: {1}})

    report = importer.load(db, [str(path)], commit=False)

    assert report["groups_not_found"] == ["НЕТ"]
    assert db.query(DaySubmission).filter_by(study_group_id=curator_group.id).one().first_period is None


def _write_landscape(path, group_code, lessons_by_day):
    """Вторая форма выгрузки: дни по строкам, группы по колонкам (блок — 4 колонки)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01.09.2026-06.09.2026"
    for _ in range(3):
        ws.append([])
    ws.append(["День недели", "Занятие", "Урок", "Дисциплина"])
    ws.append([])
    ws.append([None, "№", "№", group_code])
    ws.append([])
    for day in sorted(lessons_by_day):
        for lesson in range(1, 13):
            row = [None, None, str(lesson), None]
            if lesson == 1:
                row[0] = f"{day:%d.%m.%Y}  ВТОР"
            if (lesson + 1) // 2 in lessons_by_day[day]:
                row[3] = "Дисциплина"
            ws.append(row)
    wb.save(path)


def test_landscape_layout_and_year_suffix(tmp_path):
    p1, p2 = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _write_landscape(p1, "ИИ112", {D1: {3}, D2: {1, 2}})
    _write_schedule(p2, "ИИ112-26", {D1: {2}})

    assert importer.read_schedule(str(p1)) == {("ИИ112", D1): 3, ("ИИ112", D2): 1}
    assert importer.read_schedule(str(p2)) == {("ИИ112", D1): 2}


def test_non_schedule_xlsx_is_skipped(db, tmp_path):
    other = tmp_path / "other.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["не расписание"])
    wb.save(other)

    report = importer.load(db, [str(other)], commit=False)

    assert report["skipped_files"] == ["other.xlsx"] and report["updated"] == 0
