"""Загрузка выгрузки электронного журнала преподавателя: разбор формата, права, сверка. Только синтетические ФИО."""
import datetime
import io

import openpyxl
import pytest

from app.core.time import today_local
from app.models import AttendanceMark, LessonMark, MarkCode, Student, StudyGroup
from app.services import journal_import_service as svc

XLSX = "application/octet-stream"


def build_journal(names, group_token="ТТ1111", code="МДК.99.01", subject="Тестовые сети", page2=True, extra=None) -> bytes:
    """Выгрузка в том же устройстве, что и у колледжа: страницы по 50 строк, объединённые шапки, подпись справа."""
    extra = extra or {}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "МДК9901 Тестовые сет123456"

    def page(top, header_days, month_blocks, rows, students, period, lessons):
        ws.cell(top, 1, "№"); ws.cell(top, 2, "Студент")
        ws.merge_cells(start_row=top, start_column=1, end_row=top + 1, end_column=1)
        ws.merge_cells(start_row=top, start_column=2, end_row=top + 1, end_column=2)
        for first, last, month in month_blocks:
            ws.cell(top, first, month)
            ws.merge_cells(start_row=top, start_column=first, end_row=top, end_column=last)
        for col, day in header_days.items():
            ws.cell(top + 1, col, day)
        date_col = 21
        ws.cell(top, date_col, "Дата"); ws.cell(top, date_col + 1, "Тема"); ws.cell(top, date_col + 2, "Домашнее задание")
        for i, name in enumerate(students):
            r = top + 2 + i
            ws.cell(r, 1, str(i + 1)); ws.cell(r, 2, name + "  ")
            for col in range(3, 20):
                ws.cell(r, col, rows.get((i, col), ""))
        for i, (d, topic, hw) in enumerate(lessons):
            ws.cell(top + 2 + i, date_col, d); ws.cell(top + 2 + i, date_col + 1, topic); ws.cell(top + 2 + i, date_col + 2, hw)
        ws.cell(top + 40, date_col, f"{code} {group_token} группа, {subject}")
        ws.cell(top + 42, date_col, f"Преподаватель: Тестов Т.Т. {period}")

    # страница 1: 01.09 — пять столбцов (E:F под одним объединённым числом), 08.09 — два
    days1 = {3: "1", 4: "1", 5: "1", 7: "1", 8: "8", 9: "8"}
    page(1, days1, [(3, 19, "сен")], {(0, 3): "5", (0, 4): "н", (1, 6): "н", (1, 8): "4", (2, 9): "ку"} | extra.get("p1", {}), names, "01.09.2026-22.09.2026",
         [("01.09", "Введение", "не задано"), ("01.09", "Введение", "не задано"), ("01.09", "Кабели", "стр. 5"), ("01.09", "Кабели", "стр. 5"), ("08.09", "Адреса", "не задано")])
    ws.merge_cells(start_row=2, start_column=5, end_row=2, end_column=6)
    if page2:
        page(51, {3: "29", 4: "29", 5: "6", 6: "6", 7: "3", 8: "3"}, [(3, 4, "сен"), (5, 6, "окт"), (7, 8, "ноя")],
             {(0, 3): "3", (1, 5): "н", (1, 6): "н"} | extra.get("p2", {}), names, "29.09.2026-10.11.2026", [("29.09", "VLAN", "не задано")])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def parse(content):
    return svc.parse_workbook(content, datetime.date(2026, 10, 29))


def group_names(db, group, n=3):
    return [s.full_name for s in db.query(Student).filter(Student.study_group_id == group.id).order_by(Student.id).limit(n)]


def test_parse_dates_slots_and_values():
    parsed = parse(build_journal(["Иванов Иван Иванович", "Петрова Анна Сергеевна", "Сидоров Пётр Петрович"]))
    assert (parsed.subject_code, parsed.subject_name) == ("МДК.99.01", "Тестовые сети")
    assert parsed.group_tokens == ["ТТ1111"]
    assert parsed.period == (datetime.date(2026, 9, 1), datetime.date(2026, 11, 10))
    ivan = parsed.students[0]
    # E и F объединены: у F та же дата, что у E (01.09), а слот — следующий по порядку
    assert {(c.date, c.slot) for c in ivan.cells if c.absent} == {(datetime.date(2026, 9, 1), 2)}
    assert [(c.date, c.slot, c.grade) for c in ivan.cells if c.grade] == [
        (datetime.date(2026, 9, 1), 1, 5), (datetime.date(2026, 9, 29), 1, 3),
    ]
    petr = parsed.students[1]
    assert {(c.date, c.slot) for c in petr.cells if c.absent} == {
        (datetime.date(2026, 9, 1), 4), (datetime.date(2026, 10, 6), 1), (datetime.date(2026, 10, 6), 2),
    }
    # неизвестное значение не теряется молча
    assert parsed.unknown_values == {"ку": 1}
    assert datetime.date(2026, 11, 3) in parsed.dates
    by_slot = {(l.date, l.slot): l for l in parsed.lessons}
    assert by_slot[(datetime.date(2026, 9, 1), 1)].topic == "Введение"
    assert by_slot[(datetime.date(2026, 9, 1), 3)].homework == "стр. 5"
    assert by_slot[(datetime.date(2026, 9, 1), 5)].topic is None  # пятый столбец дня без строки в списке уроков
    assert by_slot[(datetime.date(2026, 9, 8), 1)].homework is None  # «не задано» не хранится
    assert by_slot[(datetime.date(2026, 9, 29), 1)].topic == "VLAN"


def test_not_a_journal():
    wb = openpyxl.Workbook()
    wb.active.append(["что-то", "другое"])
    out = io.BytesIO(); wb.save(out)
    with pytest.raises(svc.ImportFileError):
        parse(out.getvalue())
    with pytest.raises(svc.ImportFileError):
        parse(b"not an xlsx")


def test_apply_preview_and_replace(client, db, curator_headers, curator_group):
    names = group_names(db, curator_group)
    content = build_journal(names, group_token=curator_group.code)

    preview = client.post("/journal/preview", content=content, headers=curator_headers)
    assert preview.status_code == 200, preview.text
    assert preview.json()["applied"] is False and preview.json()["counts"]["students_matched"] == 3
    assert db.query(LessonMark).count() == 0

    applied = client.post("/journal/apply", content=content, headers=curator_headers).json()
    assert applied["applied"] is True
    first = db.query(LessonMark).count()
    assert first == applied["counts"]["grades_new"] + applied["counts"]["absences_new"] > 0
    assert any("ку" in w for w in applied["warnings"])

    # повторная загрузка — без дублей
    again = client.post("/journal/apply", content=content, headers=curator_headers).json()
    assert again["counts"]["unchanged"] == first and again["counts"]["grades_new"] == 0
    assert db.query(LessonMark).count() == first

    # преподаватель исправил оценку и стёр «н»: журнал главный
    fixed = build_journal(names, group_token=curator_group.code, extra={"p1": {(0, 3): "4", (0, 4): ""}})
    changed = client.post("/journal/apply", content=fixed, headers=curator_headers).json()
    assert changed["counts"]["updated"] == 1 and changed["counts"]["removed"] == 1
    db.expire_all()
    assert db.query(LessonMark).count() == first - 1


def test_other_group_and_unknown_students_skipped(client, db, curator_headers, curator_group, db_second_curator):
    mine = group_names(db, curator_group, 2)
    other_group = db.query(StudyGroup).filter(StudyGroup.id != curator_group.id).order_by(StudyGroup.id).first()
    db.add(Student(last_name="Чужаков", first_name="Тест", middle_name="Тестович", study_group_id=other_group.id,
                   enrolled_at=datetime.date(2026, 9, 1)))
    db.commit()
    foreign = "Чужаков Тест Тестович"
    content = build_journal(mine + [foreign, "Несуществующий Студент Тестович"], group_token=curator_group.code)
    report = client.post("/journal/apply", content=content, headers=curator_headers).json()
    assert report["counts"]["students_matched"] == 2
    assert report["out_of_scope"] == [foreign]
    assert report["not_found"] == ["Несуществующий Студент Тестович"]
    ids = {m.student_id for m in db.query(LessonMark)}
    assert len(ids) <= 2


def test_blocked_when_nobody_matches(client, curator_headers):
    report = client.post("/journal/apply", content=build_journal(["Нет Такого Человека"]), headers=curator_headers).json()
    assert report["blocked"] and report["applied"] is False


def test_bad_file_is_400(client, curator_headers):
    assert client.post("/journal/preview", content=b"nope", headers=curator_headers).status_code == 400
    assert client.post("/journal/preview", content=b"", headers=curator_headers).status_code == 400


def test_requires_login(client):
    assert client.post("/journal/preview", content=b"x").status_code == 401


def test_grades_and_discrepancies_with_scope(client, db, curator_headers, curator_group, db_second_curator, curator_user):
    names = group_names(db, curator_group)
    client.post("/journal/apply", content=build_journal(names, group_token=curator_group.code), headers=curator_headers)

    grades = client.get(f"/journal/groups/{curator_group.id}/grades", headers=curator_headers).json()
    assert [s["code"] for s in grades["subjects"]] == ["МДК.99.01"]
    first = next(s for s in grades["students"] if s["full_name"] == names[0])
    stat = first["subjects"][0]
    assert stat["count"] == 2 and stat["average"] == 4.0 and stat["absent_hours"] == 1 and stat["low_average"] is False

    # сверка: у первого студента «н» 01.09; куратор отметил присутствие — расхождение, у второго отметки нет
    present = db.query(MarkCode).filter(MarkCode.counts_as_present.is_(True)).first()
    student = next(s for s in db.query(Student).filter(Student.study_group_id == curator_group.id) if s.full_name == names[0])
    db.add(AttendanceMark(student_id=student.id, date=datetime.date(2026, 9, 1), mark_code_id=present.id, created_by_user_id=curator_user.id))
    db.commit()
    rows = client.get(
        f"/journal/groups/{curator_group.id}/discrepancies?date_from=2026-09-01&date_to=2026-11-30", headers=curator_headers,
    ).json()
    by = {(r["full_name"], r["date"]): r for r in rows}
    assert by[(names[0], "2026-09-01")]["kind"] == "present"
    assert by[(names[0], "2026-09-01")]["curator_code"] == present.code
    assert by[(names[1], "2026-09-01")]["kind"] == "no_mark"

    # чужой куратор группу не видит
    other = client.post("/auth/login", json={"username": db_second_curator.username, "password": "SecondCura#58x!"})
    headers = {"Authorization": f"Bearer {other.json()['access_token']}"} if "access_token" in other.json() else dict(other.headers)
    assert client.get(f"/journal/groups/{curator_group.id}/grades", headers=headers).status_code in (401, 403)


def test_admin_sees_any_group_and_period_limits(client, admin_headers, curator_group):
    assert client.get(f"/journal/groups/{curator_group.id}/grades", headers=admin_headers).status_code == 200
    bad = client.get(f"/journal/groups/{curator_group.id}/discrepancies?date_from=2026-10-10&date_to=2026-10-01", headers=admin_headers)
    assert bad.status_code == 400


def test_journal_view_grid(client, db, curator_headers, curator_group, db_second_curator):
    names = group_names(db, curator_group)
    assert client.get(f"/journal/groups/{curator_group.id}/subjects", headers=curator_headers).json() == []
    client.post("/journal/apply", content=build_journal(names, group_token=curator_group.code), headers=curator_headers)

    subjects = client.get(f"/journal/groups/{curator_group.id}/subjects", headers=curator_headers).json()
    assert [(s["code"], s["name"], s["lessons"]) for s in subjects] == [("МДК.99.01", "Тестовые сети", 7 + 6)]
    grid = client.get(f"/journal/groups/{curator_group.id}/subjects/МДК.99.01", headers=curator_headers).json()
    keys = [l["key"] for l in grid["lessons"]]
    assert keys[0] == "2026-09-01#1" and keys == sorted(keys, key=lambda k: (k.split("#")[0], int(k.split("#")[1])))
    assert grid["lessons"][0]["topic"] == "Введение"
    first = next(s for s in grid["students"] if s["full_name"] == names[0])
    assert first["cells"]["2026-09-01#1"] == "5" and first["cells"]["2026-09-01#2"] == "н"
    assert first["average"] == 4.0 and first["absent_hours"] == 1
    assert first["low_average"] is False and first["twos"] == 0
    # пустая ячейка — не хранится: присутствовал, оценки нет
    assert "2026-09-01#3" not in first["cells"]

    ranged = client.get(
        f"/journal/groups/{curator_group.id}/subjects/МДК.99.01?date_from=2026-09-29&date_to=2026-09-30", headers=curator_headers,
    ).json()
    assert {l["date"] for l in ranged["lessons"]} == {"2026-09-29"}
    assert client.get(f"/journal/groups/{curator_group.id}/subjects/НЕТ.00", headers=curator_headers).status_code == 404

    # куратор другой группы журнал не открывает
    login = client.post("/auth/login", json={"username": db_second_curator.username, "password": "SecondCura#58x!"})
    other = {"Authorization": f"Bearer {login.json()['access_token']}"}
    assert client.get(f"/journal/groups/{curator_group.id}/subjects", headers=other).status_code == 403


def test_low_average_and_debt_flags(client, db, curator_headers, curator_group):
    """Средняя ниже 4 подсвечивается куратору, «задолженность» — оценка 2 (решение Denis 29.10.2026)."""
    names = group_names(db, curator_group)
    extra = {"p1": {(1, 3): "3", (1, 4): "2", (2, 3): "4", (2, 4): "4"}}
    client.post("/journal/apply", content=build_journal(names, group_token=curator_group.code, extra=extra), headers=curator_headers)
    grid = client.get(f"/journal/groups/{curator_group.id}/subjects/МДК.99.01", headers=curator_headers).json()
    by = {s["full_name"]: s for s in grid["students"]}
    assert by[names[1]]["average"] == 3.0 and by[names[1]]["low_average"] is True and by[names[1]]["twos"] == 1
    assert by[names[2]]["average"] == 4.0 and by[names[2]]["low_average"] is False
    assert by[names[0]]["low_average"] is False
    grades = client.get(f"/journal/groups/{curator_group.id}/grades", headers=curator_headers).json()
    weak = {s["full_name"] for s in grades["students"] if any(x["low_average"] for x in s["subjects"])}
    assert weak == {names[1]}
