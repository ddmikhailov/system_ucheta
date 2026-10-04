"""Пол студента в досье: хранение, импорт из Excel (в том числе старым шаблоном), склонение документов."""
import datetime
import io
import re
import zipfile

import pytest
from openpyxl import Workbook

from app.models import AuditLog, Student, StudentProfile
from app.services import absence_sheet_service as sheet
from app.services import calendar_service
from app.services import dossier_import_service as imp

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_T = r"<w:t(?: [^>]*)?>(.*?)</w:t>"


def _first_student(db, group):
    return db.query(Student).filter(Student.study_group_id == group.id).order_by(Student.id).first()


def _profile_url(sid):
    return f"/students/{sid}/dossier/profile"


# ---------- хранение ----------

def test_gender_is_saved_read_and_cleared(client, curator_headers, curator_group, imported, db):
    s = _first_student(db, curator_group)
    assert client.get(f"/students/{s.id}/dossier", headers=curator_headers).json()["profile"]["gender"] is None
    for value in ("female", "male"):
        r = client.put(_profile_url(s.id), headers=curator_headers, json={"gender": value})
        assert r.status_code == 200 and r.json()["profile"]["gender"] == value
        assert client.get(f"/students/{s.id}/dossier", headers=curator_headers).json()["profile"]["gender"] == value
    db.expire_all()
    assert db.get(StudentProfile, s.id).gender == "male"
    cleared = client.put(_profile_url(s.id), headers=curator_headers, json={"gender": None})
    assert cleared.json()["profile"]["gender"] is None


def test_unknown_gender_value_is_rejected(client, curator_headers, curator_group, imported, db):
    s = _first_student(db, curator_group)
    assert client.put(_profile_url(s.id), headers=curator_headers, json={"gender": "other"}).status_code == 422
    assert client.put(_profile_url(s.id), headers=curator_headers, json={"gender": "м"}).status_code == 422


def test_gender_change_is_audited_without_the_value(client, curator_headers, curator_group, imported, db):
    s = _first_student(db, curator_group)
    client.put(_profile_url(s.id), headers=curator_headers, json={"gender": "female"})
    entry = db.query(AuditLog).filter(AuditLog.action == "dossier.profile_update").one()
    assert entry.new_value == "gender" and "female" not in (entry.new_value or "")


def test_other_curator_cannot_set_gender(client, curator_group, db_second_curator, imported, db):
    from tests.conftest import _login
    s = _first_student(db, curator_group)
    other = _login(client, db_second_curator.username, "SecondCurator123!")
    assert client.put(_profile_url(s.id), headers=other, json={"gender": "male"}).status_code == 403


# ---------- импорт из Excel ----------

def _file(rows, legacy=False):
    headers = imp._headers(legacy=legacy)
    wb = Workbook()
    ws = wb.active
    ws.title = "Досье"
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _who(s):
    return {"Группа": s.study_group.code, "Фамилия": s.last_name, "Имя": s.first_name, "Отчество": s.middle_name}


def _post(client, path, headers, content):
    return client.post(path, headers={**headers, "Content-Type": XLSX}, content=content)


def test_template_contains_the_gender_column_and_hint(client, admin_headers, imported):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get("/dossier-import/template", headers=admin_headers).content))
    headers = [c.value for c in wb.worksheets[0][1]]
    assert "Пол (м/ж)" in headers and headers.index("Пол (м/ж)") == headers.index("Дата рождения") + 1
    assert any("Пол: м или ж" in (row[0] or "") for row in wb["Инструкция"].iter_rows(values_only=True))


@pytest.mark.parametrize("cell,expected", [("м", "male"), ("Ж", "female"), ("мужской", "male"), (" женский ", "female")])
def test_import_understands_gender_spellings(client, admin_headers, imported, db, cell, expected):
    s = db.query(Student).first()
    r = _post(client, "/dossier-import/apply", admin_headers, _file([{**_who(s), "Пол (м/ж)": cell}]))
    assert r.status_code == 200 and r.json()["updated"] == 1
    db.expire_all()
    assert db.get(StudentProfile, s.id).gender == expected


def test_import_reports_a_bad_gender_and_empty_cell_keeps_the_value(client, admin_headers, imported, db):
    s = db.query(Student).first()
    _post(client, "/dossier-import/apply", admin_headers, _file([{**_who(s), "Пол (м/ж)": "ж"}]))
    preview = _post(client, "/dossier-import/preview", admin_headers, _file([{**_who(s), "Пол (м/ж)": "оно"}])).json()
    assert preview["with_errors"] == 1
    assert any("Пол (м/ж)" in e and "ожидалось м или ж" in e for e in preview["rows"][0]["errors"])
    _post(client, "/dossier-import/apply", admin_headers, _file([{**_who(s), "Телефон": "+7 900 000-00-00"}]))
    db.expire_all()
    profile = db.get(StudentProfile, s.id)
    assert profile.gender == "female" and profile.phone == "+7 900 000-00-00"  # пустая ячейка пол не стёрла


def test_files_from_the_first_template_version_are_still_accepted(client, admin_headers, imported, db):
    """Шаблоны, розданные до появления колонки «Пол», продолжают работать (пол из них не загружается)."""
    s = db.query(Student).first()
    legacy = _file([{**_who(s), "Дата рождения": "05.03.2008", "Телефон": "+7 900 111-22-33", "Сирота (да/нет)": "да"}], legacy=True)
    assert "Пол (м/ж)" not in imp._headers(legacy=True) and len(imp._headers(legacy=True)) == len(imp._headers()) - 1
    preview = _post(client, "/dossier-import/preview", admin_headers, legacy).json()
    assert (preview["total"], preview["ready"], preview["with_errors"]) == (1, 1, 0)
    assert _post(client, "/dossier-import/apply", admin_headers, legacy).json()["updated"] == 1
    db.expire_all()
    profile = db.get(StudentProfile, s.id)
    assert profile.birth_date == datetime.date(2008, 3, 5) and profile.phone == "+7 900 111-22-33" and profile.gender is None


def test_a_file_with_foreign_headers_is_still_rejected(client, admin_headers):
    wb = Workbook()
    wb.active.title = "Досье"
    wb.active.append(["Группа", "Фамилия", "Имя", "Отчество", "Что-то другое"])
    buf = io.BytesIO()
    wb.save(buf)
    r = _post(client, "/dossier-import/preview", admin_headers, buf.getvalue())
    assert r.status_code == 400 and "Заголовки не совпадают" in r.json()["detail"]


# ---------- документы ----------

D = datetime.date
ROWS = [sheet.SheetRow(D(2026, 10, 1), "Неуважительная причина", False, False)]


def _doc(gender):
    data = sheet.build_docx(student_name="Лебедева Елена Андреевна", group_code="СА172", department_name="Диджитал",
                            curator_name="Иванова Анна Ивановна", date_from=D(2026, 10, 1), date_to=D(2026, 10, 7),
                            rows=ROWS, gender=gender)
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()
    return "\n".join("".join(re.findall(_T, p, flags=re.S)) for p in re.findall(r"<w:p>.*?</w:p>|<w:p .*?</w:p>", xml, flags=re.S))


def test_female_forms():
    text = _doc("female")
    for phrase in ("от обучающейся группы СА172", "обучающаяся группы СА172, подтверждаю, что ознакомлена со сведениями",
                   "С Правилами внутреннего распорядка обучающихся ГБПОУ КАИТ № 20 ознакомлена.", "согласна / имею замечания",
                   "(Ф.И.О. обучающейся полностью)", "Подпись обучающейся"):
        assert phrase in text, phrase
    assert "(а)" not in text and "(на)" not in text and "(аяся)" not in text and "(ейся)" not in text


def test_male_forms():
    text = _doc("male")
    for phrase in ("от обучающегося группы СА172", "обучающийся группы СА172, подтверждаю, что ознакомлен со сведениями",
                   "ознакомлен. Мне разъяснены", "согласен / имею замечания", "(Ф.И.О. обучающегося полностью)", "Подпись обучающегося"):
        assert phrase in text, phrase
    assert "(а)" not in text and "(на)" not in text and "(аяся)" not in text and "(ейся)" not in text


@pytest.mark.parametrize("gender", [None, "", "unknown"])
def test_unknown_gender_keeps_the_neutral_forms_of_the_template(gender):
    text = _doc(gender)
    for phrase in ("обучающегося(ейся) группы СА172", "обучающийся(аяся) группы СА172", "ознакомлен(а) со сведениями", "согласен(на) / имею замечания"):
        assert phrase in text, phrase


def test_rules_quoted_from_the_local_act_are_never_declined():
    for gender in ("male", "female"):
        text = _doc(gender)
        assert "обучающийся обязан посещать учебные занятия по расписанию" in text
        assert "обучающийся должен приходить в Колледж" in text


def _mark(client, headers, group, student, code, day):
    r = client.post(f"/curator/groups/{group.id}/day/submit?date={day}", headers=headers, json={
        "exceptions": [{"student_id": student.id, "mark_code": code, "comment": None, "basis_reference": None}]})
    assert r.status_code == 200, r.text


@pytest.mark.parametrize("gender,forms", [("female", ("обучающаяся", "ознакомлена", "студентки")), ("male", ("обучающийся", "ознакомлен", "студента")),
                                          (None, ("обучающийся(аяся)", "ознакомлен(а)", "студента"))])
def test_sheet_and_parents_message_follow_the_saved_gender(client, curator_headers, curator_group, imported, db, today, gender, forms):
    s = _first_student(db, curator_group)
    day = calendar_service.study_days_between(
        db, today - datetime.timedelta(days=10), today, study_group_id=curator_group.id, course=curator_group.course)[-1]
    _mark(client, curator_headers, curator_group, s, "н", day)
    if gender:
        client.put(_profile_url(s.id), headers=curator_headers, json={"gender": gender})
    r = client.get(f"/students/{s.id}/absence-sheet?date_from={day}&date_to={day}", headers=curator_headers)
    xml = zipfile.ZipFile(io.BytesIO(r.content)).read("word/document.xml").decode()
    assert forms[0] in xml and forms[1] in xml
    message = client.get(f"/students/{s.id}/absence-message", headers=curator_headers).json()["text"]
    assert f"о посещаемости {forms[2]}: " in message
