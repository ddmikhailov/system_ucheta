import io

import openpyxl


def test_excel_export_contains_all_groups(client, admin_headers, imported, today):
    r = client.get(f"/export/excel?date_from={today}&date_to={today}", headers=admin_headers)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")

    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "ИТОГ" in wb.sheetnames
    itog = wb["ИТОГ"]
    # 44 группы + строка заголовка.
    assert itog.max_row == 45
    assert "ИИ112" in wb.sheetnames


def test_excel_export_forbidden_for_curator(client, curator_headers, today):
    r = client.get(f"/export/excel?date_from={today}&date_to={today}", headers=curator_headers)
    assert r.status_code == 403


def test_pdf_export_for_own_group(client, curator_headers, curator_group, today):
    r = client.get(
        f"/export/pdf/{curator_group.id}?date_from={today}&date_to={today}", headers=curator_headers
    )
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content[:4] == b"%PDF"


def test_pdf_export_forbidden_for_other_groups_curator(client, curator_headers, db, curator_group, today):
    from app.models import StudyGroup

    other_group = db.query(StudyGroup).filter(StudyGroup.id != curator_group.id).first()
    r = client.get(
        f"/export/pdf/{other_group.id}?date_from={today}&date_to={today}", headers=curator_headers
    )
    assert r.status_code == 403
