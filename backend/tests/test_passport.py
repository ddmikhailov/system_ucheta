"""Социальный паспорт группы: подсчёты из досье, доступ по ролям, журнал, Excel."""
import datetime
import io

from openpyxl import load_workbook

from app.core.config import get_settings
from app.models import Department, DossierAccessLog, Student, StudyGroup


def _fill(client, headers, student, **payload):
    r = client.put(f"/students/{student.id}/dossier/profile", headers=headers, json=payload)
    assert r.status_code == 200, r.text


def _group_students(db, group):
    return db.query(Student).filter(Student.study_group_id == group.id).order_by(Student.id).all()


def _category(data, key):
    return next(c for c in data["categories"] if c["key"] == key)


def test_group_passport_counts_categories_and_dossier_completeness(client, curator_headers, curator_group, db):
    a, b, c = _group_students(db, curator_group)[:3]
    adult = (datetime.date.today() - datetime.timedelta(days=365 * 20)).isoformat()
    minor = (datetime.date.today() - datetime.timedelta(days=365 * 16)).isoformat()
    _fill(client, curator_headers, a, birth_date=minor, funding="budget",
          special={"is_orphan": True, "disability_group": "3", "scholarship": "соц."})
    _fill(client, curator_headers, b, birth_date=adult, funding="contract",
          special={"large_family": True, "pdn_kdn": True})
    client.post(f"/students/{a.id}/dossier/guardians", headers=curator_headers,
                json={"full_name": "Мать А.", "relation": "мать"})

    r = client.get(f"/passport/group/{curator_group.id}", headers=curator_headers)
    assert r.status_code == 200, r.text
    data = r.json()
    total = len(_group_students(db, curator_group))
    assert data["students_total"] == total
    assert (data["minors"], data["adults"], data["budget"], data["contract"]) == (1, 1, 1, 1)
    assert data["birth_date_missing"] == total - 2
    assert data["funding_missing"] == total - 2
    assert data["no_guardians"] == total - 1
    assert data["dossier_empty"] == total - 2  # а и b заполнены (a ещё и с представителем)
    assert _category(data, "orphan")["count"] == 1 and _category(data, "orphan")["names"] == [a.full_name]
    assert _category(data, "disability")["count"] == 1
    assert _category(data, "scholarship")["count"] == 1
    assert _category(data, "large_family")["names"] == [b.full_name]
    assert _category(data, "ovz")["count"] == 0
    assert data["special_available"] is True


def test_curator_sees_only_own_groups_passport(client, curator_headers, curator_group, db):
    other = db.query(StudyGroup).filter(StudyGroup.id != curator_group.id).first()
    assert client.get(f"/passport/group/{other.id}", headers=curator_headers).status_code == 403
    assert client.get("/passport/group/999999", headers=curator_headers).status_code == 404
    rows = client.get("/passport/summary", headers=curator_headers).json()["rows"]
    assert [r["group_id"] for r in rows] == [curator_group.id]
    assert client.get(f"/passport/export?group_id={other.id}", headers=curator_headers).status_code == 403


def test_summary_scoped_by_role(client, admin_headers, dept_head_headers, dept_head_user, imported, db):
    other = Department(name="Чужое")
    db.add(other)
    db.flush()
    db.add(StudyGroup(code="OTH-1", course=1, department_id=other.id))
    db.commit()
    everything = client.get("/passport/summary", headers=admin_headers).json()
    mine = client.get("/passport/summary", headers=dept_head_headers).json()
    assert any(r["group_code"] == "OTH-1" for r in everything["rows"])
    assert all(r["group_code"] != "OTH-1" for r in mine["rows"])
    assert mine["totals"]["students_total"] == sum(r["students_total"] for r in mine["rows"])
    only_other = client.get(f"/passport/summary?department_id={other.id}", headers=admin_headers).json()
    assert [r["group_code"] for r in only_other["rows"]] == ["OTH-1"]
    assert [c["key"] for c in everything["categories"]][:2] == ["orphan", "guardianship"]


def test_named_passport_is_logged_in_dossier_access_log(client, curator_headers, curator_group, admin_headers, db):
    a, b = _group_students(db, curator_group)[:2]
    _fill(client, curator_headers, a, special={"is_orphan": True})
    _fill(client, curator_headers, b, special={"low_income": True})
    before = db.query(DossierAccessLog).count()
    client.get(f"/passport/group/{curator_group.id}", headers=curator_headers)
    db.expire_all()
    assert db.query(DossierAccessLog).count() == before + 2
    logged = client.get(f"/students/{a.id}/dossier/access-log", headers=admin_headers).json()
    assert any(row["included_special"] for row in logged)


def test_passport_without_encryption_key_hides_special_counts(client, curator_headers, curator_group, db, monkeypatch):
    s = _group_students(db, curator_group)[0]
    _fill(client, curator_headers, s, funding="budget", special={"is_orphan": True})
    monkeypatch.setattr(get_settings(), "dossier_encryption_key", "")
    data = client.get(f"/passport/group/{curator_group.id}", headers=curator_headers).json()
    assert data["special_available"] is False
    assert all(c["count"] is None and c["names"] == [] for c in data["categories"])
    assert data["budget"] == 1  # обычные показатели считаются и без ключа
    summary = client.get("/passport/summary", headers=curator_headers).json()
    assert summary["special_available"] is False and summary["totals"]["counts"]["orphan"] is None


def test_export_excel_summary_and_group_lists(client, curator_headers, curator_group, admin_headers, db):
    a = _group_students(db, curator_group)[0]
    _fill(client, curator_headers, a, special={"is_orphan": True})

    r = client.get("/passport/export", headers=admin_headers)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Сводка"]  # поимённых списков в общей выгрузке нет
    rows = list(wb["Сводка"].iter_rows(values_only=True))
    assert rows[0][0] == "Группа" and rows[-1][0] == "Итого"

    r = client.get(f"/passport/export?group_id={curator_group.id}", headers=curator_headers)
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Сводка", "Списки"]
    assert ["Сироты", a.full_name] in [list(row) for row in wb["Списки"].iter_rows(min_row=2, values_only=True)]


def test_staff_and_edu_department_see_all_groups(client, edu_department_headers, db, imported):
    data = client.get("/passport/summary", headers=edu_department_headers).json()
    assert len(data["rows"]) == db.query(StudyGroup).filter(StudyGroup.is_active.is_(True)).count()
