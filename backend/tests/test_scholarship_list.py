"""Список группы на стипендии (п. 5.1.14): виды из досье, Excel и Word, доступ по ролям, журнал просмотров. Данные выдуманы."""
import io
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook

from app.core.config import get_settings
from app.models import AuditLog, Department, DossierAccessLog, Student, StudyGroup

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _students(db, group):
    return db.query(Student).filter(Student.study_group_id == group.id).order_by(Student.last_name, Student.first_name).all()


def _fill(client, headers, student, **special):
    r = client.put(f"/students/{student.id}/dossier/profile", headers=headers, json={"special": special})
    assert r.status_code == 200, r.text


def _get(client, headers, group, fmt):
    return client.get(f"/passport/group/{group.id}/scholarships?format={fmt}", headers=headers)


def _xlsx_rows(content):
    ws = load_workbook(io.BytesIO(content)).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _docx_cells(content):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert archive.testzip() is None
        root = ET.fromstring(archive.read("word/document.xml"))
    return [["".join(t.text or "" for t in tc.iter(f"{W}t")) for tc in tr.findall(f"{W}tc")] for tr in root.iter(f"{W}tr")]


def test_students_are_grouped_by_kind_in_both_formats(client, curator_headers, curator_group, db):
    a, b, c, d = _students(db, curator_group)[:4]
    _fill(client, curator_headers, a, scholarship="Социальная")
    _fill(client, curator_headers, b, scholarship="  академическая ")
    _fill(client, curator_headers, c, scholarship="социальная")  # тот же вид: регистр и пробелы не различают
    _fill(client, curator_headers, d, large_family=True)  # без стипендии в список не попадает

    rows = _xlsx_rows(_get(client, curator_headers, curator_group, "xlsx").content)
    table = [r for r in rows if isinstance(r[0], int) and r[1] and r[1] != "Итого" and r[2] in ("Социальная", "академическая")]
    social = [r for r in table if r[2] == "Социальная"]
    assert [r[1] for r in social] == sorted([a.full_name, c.full_name], key=lambda n: n)  # по алфавиту внутри вида
    assert {r[1] for r in table} == {a.full_name, b.full_name, c.full_name}
    assert [r[0] for r in table] == [1, 2, 3]  # социальная (2 студента) идёт раньше академической (1)
    assert all(d.full_name != r[1] for r in rows)

    cells = _docx_cells(_get(client, curator_headers, curator_group, "docx").content)
    assert cells[0] == ["№", "ФИО", "Вид стипендии"]
    assert [r[1] for r in cells[1:]] == [r[1] for r in table]
    assert d.full_name not in str(cells)


def test_empty_list_still_builds_a_document(client, curator_headers, curator_group):
    xlsx = _xlsx_rows(_get(client, curator_headers, curator_group, "xlsx").content)
    assert any("нет студентов со стипендией" in str(r) for r in xlsx)
    docx = _docx_cells(_get(client, curator_headers, curator_group, "docx").content)
    assert docx[-1][1] == "В группе нет студентов со стипендией"


def test_text_from_dossier_cannot_become_a_formula(client, curator_headers, curator_group, db):
    a = _students(db, curator_group)[0]
    _fill(client, curator_headers, a, scholarship="=1+1")
    content = _get(client, curator_headers, curator_group, "xlsx").content
    cell = next(c for row in load_workbook(io.BytesIO(content)).active.iter_rows() for c in row if c.value == "=1+1")
    assert cell.data_type == "s"


def test_access_follows_the_passport_rules(
    client, curator_headers, admin_headers, dept_head_headers, edu_department_headers, curator_group, db,
):
    other_dept = Department(name="Чужое")
    db.add(other_dept)
    db.flush()
    foreign = StudyGroup(code="OTH-1", course=1, department_id=other_dept.id)
    db.add(foreign)
    db.commit()
    mine, own_other = curator_group, db.query(StudyGroup).filter(StudyGroup.id != curator_group.id, StudyGroup.code != "OTH-1").first()

    assert client.get(f"/passport/group/{mine.id}/scholarships").status_code == 401
    assert _get(client, curator_headers, mine, "xlsx").status_code == 200
    assert _get(client, curator_headers, foreign, "xlsx").status_code == 403  # куратор — только свои группы
    if own_other is not None:
        assert _get(client, curator_headers, own_other, "docx").status_code == 403
    assert _get(client, dept_head_headers, mine, "docx").status_code == 200
    assert _get(client, dept_head_headers, foreign, "docx").status_code == 403  # зав. отделением — только своё отделение
    assert _get(client, admin_headers, foreign, "xlsx").status_code == 200  # админ — весь колледж
    assert _get(client, edu_department_headers, mine, "xlsx").status_code == 200  # воспитательный отдел — как в паспорте
    assert _get(client, edu_department_headers, foreign, "xlsx").status_code == 403  # закреплён за отделением
    assert client.get("/passport/group/999999/scholarships", headers=admin_headers).status_code == 404
    assert client.get(f"/passport/group/{mine.id}/scholarships?format=pdf", headers=admin_headers).status_code == 422


def test_without_encryption_key_the_list_is_refused(client, curator_headers, curator_group, monkeypatch):
    monkeypatch.setattr(get_settings(), "dossier_encryption_key", "")
    r = _get(client, curator_headers, curator_group, "xlsx")
    assert r.status_code == 409 and "ключ шифрования" in r.json()["detail"]


def test_export_is_logged_per_student(client, curator_headers, curator_user, curator_group, db):
    a, b, c = _students(db, curator_group)[:3]
    _fill(client, curator_headers, a, scholarship="Социальная")
    _fill(client, curator_headers, b, scholarship="Академическая")
    _fill(client, curator_headers, c, large_family=True)
    db.query(DossierAccessLog).delete()
    db.commit()
    r = _get(client, curator_headers, curator_group, "docx")
    assert "filename*=UTF-8''" in r.headers["content-disposition"] and r.headers["content-disposition"].endswith(".docx")
    logged = {row.student_id for row in db.query(DossierAccessLog).filter_by(user_id=curator_user.id, included_special=True)}
    assert logged == {a.id, b.id}  # только те, чьи данные попали в список
    assert db.query(AuditLog).filter_by(action="passport.scholarships", user_id=curator_user.id).count() == 1
