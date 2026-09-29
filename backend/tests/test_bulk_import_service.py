"""Разбор Excel-таблиц посещаемости за сентябрь 2026 — см. пояснение в
app/services/bulk_import_service.py. Тесты строят синтетическую книгу той же
формы, что реальные файлы (заголовок, строки студентов, служебные строки
после них), не используют реальные ФИО."""
import io

import openpyxl

from app.services.bulk_import_service import parse_attendance_workbook


def _build_workbook(sheet_name: str, student_rows: list[tuple], trailing_rows: list[tuple] | None = None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append((sheet_name, None, "Сентябрь"))
    ws.append((None,))
    header = ["№ п/п", "ФИО"] + list(range(1, 32)) + [None, "ИУП", "Больничный лист", "Заявление", "По приказу"]
    ws.append(header)
    for row in student_rows:
        ws.append(row)
    for row in (trailing_rows or []):
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _student_row(no: int, fio: str, codes_by_day: dict[int, str]) -> tuple:
    row = [no, fio] + [None] * 31
    for day, code in codes_by_day.items():
        row[1 + day] = code
    return tuple(row)


def test_parses_marks_for_real_students():
    data = _build_workbook(
        "ИИ112",
        student_rows=[
            _student_row(1, "Тестов Иван Иванович", {5: "н", 10: "б"}),
            _student_row(2, "Тестова Мария Петровна", {}),
        ],
    )
    sheets = parse_attendance_workbook(data)
    assert len(sheets) == 1
    sheet = sheets[0]
    assert sheet.group_code == "ИИ112"
    assert sheet.error is None
    assert sheet.marks_by_student["Тестов Иван Иванович"] == {5: "н", 10: "б"}
    assert "Тестова Мария Петровна" not in sheet.marks_by_student  # ни одной отметки


def test_stops_before_trailing_summary_rows():
    data = _build_workbook(
        "ИИ112",
        student_rows=[_student_row(1, "Тестов Иван Иванович", {3: "о"})],
        trailing_rows=[
            (25, None) + (None,) * 31,  # пустая заготовленная строка без ФИО
            (None, "Начало занятий  (№ пары)") + ("Д",) + (None,) * 30,
            (None, "Количество студентов, опаздавших на занятие") + (0,) * 31,
        ],
    )
    sheets = parse_attendance_workbook(data)
    sheet = sheets[0]
    assert sheet.marks_by_student == {"Тестов Иван Иванович": {3: "о"}}


def test_unrecognized_code_reported_not_silently_dropped():
    data = _build_workbook(
        "ИИ112",
        student_rows=[_student_row(1, "Тестов Иван Иванович", {7: "эл"})],
    )
    sheets = parse_attendance_workbook(data)
    sheet = sheets[0]
    assert sheet.marks_by_student == {}
    assert sheet.unrecognized == [("Тестов Иван Иванович", 7, "эл")]


def test_missing_header_row_reports_error_not_crash():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ПУСТОЙ"
    ws.append(("что-то", "не то"))
    buf = io.BytesIO()
    wb.save(buf)
    sheets = parse_attendance_workbook(buf.getvalue())
    assert sheets[0].error is not None


def test_skips_itog_and_list1_sheets():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ИТОГ"
    ws.append(("whatever",))
    wb.create_sheet("Лист1")
    buf = io.BytesIO()
    wb.save(buf)
    sheets = parse_attendance_workbook(buf.getvalue())
    assert sheets == []
