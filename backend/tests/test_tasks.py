"""Задачи от администрации: охват, заполнение куратором, проверка, права, выгрузка."""
import datetime
import io

import pytest
from openpyxl import load_workbook

from app.core.time import utcnow
from app.models import Department, Student, StudyGroup, Task, TaskAssignment

FIELDS_STUDENT = [
    {"label": "Кружки", "type": "multiselect", "options": ["Спорт", "Танцы", "Робототехника"], "required": False},
    {"label": "Не посещает", "type": "bool", "required": False},
    {"label": "Телефон родителя", "type": "text", "required": True},
]


def _future(days=7):
    return str(datetime.date.today() + datetime.timedelta(days=days))


def _payload(**overrides):
    base = {
        "title": "Кружки доп. образования", "description": "Собрать данные", "collect_mode": "student",
        "reviewer_rule": "dept_head", "due_date": _future(), "fields": FIELDS_STUDENT,
        "scope": {"all_groups": True},
    }
    base.update(overrides)
    return base


def _create(client, headers, **overrides):
    r = client.post("/tasks", headers=headers, json=_payload(**overrides))
    assert r.status_code == 201, r.text
    return r.json()


def _my_assignment_id(client, headers, task_id):
    rows = client.get("/tasks/my", headers=headers).json()
    return next(r["id"] for r in rows if r["task_id"] == task_id)


def _students(db, group):
    return db.query(Student).filter(Student.study_group_id == group.id).order_by(Student.id).all()


def _field_keys(task):
    return {f["label"]: f["key"] for f in task["fields"]}


# ---------- создание и охват ----------

def test_admin_creates_task_for_all_groups_and_curator_is_notified(client, admin_headers, curator_headers, curator_group, imported, db):
    task = _create(client, admin_headers)
    active_groups = db.query(StudyGroup).filter(StudyGroup.is_active.is_(True)).count()
    assert task["progress"]["total"] == active_groups and task["progress"]["new"] == active_groups
    assert [f["key"] for f in task["fields"]] == ["f1", "f2", "f3"]

    mine = client.get("/tasks/my", headers=curator_headers).json()
    assert [(r["title"], r["status"], r["group_code"]) for r in mine] == [
        ("Кружки доп. образования", "new", curator_group.code)
    ]
    notifications = client.get("/notifications", headers=curator_headers).json()
    assert any(n["kind"] == "task_assigned" for n in notifications)


def test_scope_by_department_course_groups_and_exclusions(client, admin_headers, imported, db):
    groups = db.query(StudyGroup).filter(StudyGroup.is_active.is_(True)).order_by(StudyGroup.id).all()
    first, second = groups[0], groups[1]
    only_two = _create(client, admin_headers, scope={"group_ids": [first.id, second.id]})
    assert only_two["progress"]["total"] == 2
    excluded = _create(client, admin_headers, scope={"all_groups": True, "exclude_group_ids": [first.id]})
    assert excluded["progress"]["total"] == len(groups) - 1
    course = _create(client, admin_headers, scope={"courses": [first.course]})
    assert course["progress"]["total"] == sum(1 for g in groups if g.course == first.course)

    other = Department(name="Другое отделение")
    db.add(other)
    db.commit()
    r = client.post("/tasks", headers=admin_headers, json=_payload(scope={"department_ids": [other.id]}))
    assert r.status_code == 400 and "ни одна группа" in r.json()["detail"]
    assert client.post("/tasks", headers=admin_headers, json=_payload(scope={})).status_code == 400


def test_task_validation(client, admin_headers, imported):
    bad_select = [{"label": "Выбор", "type": "select", "options": []}]
    assert client.post("/tasks", headers=admin_headers, json=_payload(fields=bad_select)).status_code == 400
    assert client.post("/tasks", headers=admin_headers, json=_payload(fields=[])).status_code == 400
    past = str(datetime.date.today() - datetime.timedelta(days=1))
    assert client.post("/tasks", headers=admin_headers, json=_payload(due_date=past)).status_code == 400
    assert client.post("/tasks", headers=admin_headers, json=_payload(collect_mode="everyone")).status_code == 422


def test_dept_head_can_create_only_for_own_department(client, dept_head_headers, dept_head_user, imported, db):
    other = Department(name="Другое отделение")
    db.add(other)
    db.flush()
    foreign = StudyGroup(code="OTH-9", course=1, department_id=other.id)
    db.add(foreign)
    db.commit()
    own_count = db.query(StudyGroup).filter(
        StudyGroup.is_active.is_(True), StudyGroup.department_id == dept_head_user.department_id
    ).count()
    task = _create(client, dept_head_headers)  # «весь колледж» урезается до своего отделения
    assert task["progress"]["total"] == own_count
    assert all(a["group_code"] != "OTH-9" for a in task["assignments"])
    r = client.post("/tasks", headers=dept_head_headers, json=_payload(scope={"group_ids": [foreign.id]}))
    assert r.status_code == 400


def test_curator_and_staff_cannot_create_tasks(client, curator_headers):
    assert client.post("/tasks", headers=curator_headers, json=_payload()).status_code == 403
    assert client.get("/tasks", headers=curator_headers).status_code == 403


# ---------- жизненный цикл ----------

def test_full_lifecycle_student_mode(client, admin_headers, dept_head_headers, curator_headers, curator_group, imported, db):
    task = _create(client, admin_headers)
    aid = _my_assignment_id(client, curator_headers, task["id"])
    keys = _field_keys(task)

    detail = client.get(f"/tasks/assignments/{aid}", headers=curator_headers).json()
    students = _students(db, curator_group)
    assert len(detail["rows"]) == len(students) and detail["can_edit"] and not detail["can_review"]

    # Пока не все ответили — отправить нельзя.
    first = students[0]
    r = client.put(f"/tasks/assignments/{aid}/answers", headers=curator_headers, json={"rows": [
        {"student_id": first.id, "values": {keys["Кружки"]: ["Спорт", "Танцы"], keys["Телефон родителя"]: "+7 900"}}
    ]})
    assert r.status_code == 200 and r.json()["status"] == "in_progress"
    r = client.post(f"/tasks/assignments/{aid}/submit", headers=curator_headers)
    assert r.status_code == 400 and "нет ответа" in r.json()["detail"]

    # Обязательное поле.
    rows = [{"student_id": s.id, "values": {keys["Кружки"]: []}} for s in students]
    client.put(f"/tasks/assignments/{aid}/answers", headers=curator_headers, json={"rows": rows})
    r = client.post(f"/tasks/assignments/{aid}/submit", headers=curator_headers)
    assert r.status_code == 400 and "Телефон родителя" in r.json()["detail"]

    rows = [{"student_id": s.id, "values": {keys["Телефон родителя"]: "+7 900"}} for s in students]
    client.put(f"/tasks/assignments/{aid}/answers", headers=curator_headers, json={"rows": rows})
    r = client.post(f"/tasks/assignments/{aid}/submit", headers=curator_headers)
    assert r.status_code == 200 and r.json()["status"] == "submitted"
    assert not r.json()["can_edit"]
    assert client.put(f"/tasks/assignments/{aid}/answers", headers=curator_headers, json={"rows": rows}).status_code == 400

    # Проверяющий (зав. отделением) получил уведомление и видит очередь.
    assert any(n["kind"] == "task_submitted" for n in client.get("/notifications", headers=dept_head_headers).json())
    queue = client.get("/tasks/review-queue", headers=dept_head_headers).json()
    assert [q["id"] for q in queue] == [aid]
    assert client.get(f"/tasks/assignments/{aid}", headers=dept_head_headers).json()["can_review"] is True

    # Возврат требует комментарий; после возврата куратор снова может править.
    review_url = f"/tasks/assignments/{aid}/review"
    assert client.post(review_url, headers=dept_head_headers, json={"action": "return"}).status_code == 400
    r = client.post(review_url, headers=dept_head_headers, json={"action": "return", "comment": "У Иванова нет кружка"})
    assert r.status_code == 200 and r.json()["status"] == "returned" and r.json()["review_comment"]
    assert any(n["kind"] == "task_returned" for n in client.get("/notifications", headers=curator_headers).json())
    assert client.get(f"/tasks/assignments/{aid}", headers=curator_headers).json()["can_edit"] is True
    client.put(f"/tasks/assignments/{aid}/answers", headers=curator_headers, json={"rows": rows[:1]})
    assert client.get(f"/tasks/assignments/{aid}", headers=curator_headers).json()["status"] == "in_progress"
    assert client.post(f"/tasks/assignments/{aid}/submit", headers=curator_headers).json()["status"] == "submitted"

    r = client.post(review_url, headers=dept_head_headers, json={"action": "accept"})
    assert r.json()["status"] == "accepted"
    assert any(n["kind"] == "task_accepted" for n in client.get("/notifications", headers=curator_headers).json())
    assert client.post(review_url, headers=dept_head_headers, json={"action": "accept"}).status_code == 400


def test_group_mode_with_link_and_value_validation(client, admin_headers, curator_headers, curator_group, imported):
    fields = [
        {"label": "Видео", "type": "link", "required": True},
        {"label": "Дата съёмки", "type": "date"},
        {"label": "Участников", "type": "number"},
        {"label": "Формат", "type": "select", "options": ["Видео", "Фото"]},
    ]
    task = _create(client, admin_headers, collect_mode="group", fields=fields, title="Видеовизитка")
    aid = _my_assignment_id(client, curator_headers, task["id"])
    keys = _field_keys(task)
    url = f"/tasks/assignments/{aid}/answers"

    for bad in ({"Видео": "не ссылка"}, {"Дата съёмки": "вчера"}, {"Участников": "много"}, {"Формат": "Аудио"}):
        r = client.put(url, headers=curator_headers, json={"group_values": {keys[k]: v for k, v in bad.items()}})
        assert r.status_code == 400, bad
    assert client.put(url, headers=curator_headers, json={"group_values": {"nope": 1}}).status_code == 400
    assert client.put(url, headers=curator_headers, json={"rows": [{"student_id": 1}]}).status_code == 400

    assert client.post(f"/tasks/assignments/{aid}/submit", headers=curator_headers).status_code == 400
    r = client.put(url, headers=curator_headers, json={"group_values": {
        keys["Видео"]: "https://vk.com/video1", keys["Участников"]: "12,0", keys["Дата съёмки"]: "2026-10-01"}})
    assert r.status_code == 200 and r.json()["group_values"][keys["Участников"]] == 12
    assert client.post(f"/tasks/assignments/{aid}/submit", headers=curator_headers).json()["status"] == "submitted"


def test_selected_mode_only_included_rows_are_checked(client, admin_headers, curator_headers, curator_group, imported, db):
    fields = [{"label": "Тема", "type": "text", "required": True}]
    task = _create(client, admin_headers, collect_mode="selected", fields=fields)
    aid = _my_assignment_id(client, curator_headers, task["id"])
    key = task["fields"][0]["key"]
    a, b = _students(db, curator_group)[:2]
    detail = client.get(f"/tasks/assignments/{aid}", headers=curator_headers).json()
    assert all(r["is_included"] is False for r in detail["rows"])

    client.put(f"/tasks/assignments/{aid}/answers", headers=curator_headers, json={"rows": [
        {"student_id": a.id, "is_included": True, "values": {}}, {"student_id": b.id, "is_included": False}]})
    r = client.post(f"/tasks/assignments/{aid}/submit", headers=curator_headers)
    assert r.status_code == 400 and a.full_name in r.json()["detail"]
    client.put(f"/tasks/assignments/{aid}/answers", headers=curator_headers, json={"rows": [
        {"student_id": a.id, "is_included": True, "values": {key: "Конкурс"}}]})
    assert client.post(f"/tasks/assignments/{aid}/submit", headers=curator_headers).json()["status"] == "submitted"
    # Студент из чужой группы в ответ не попадает.
    other = db.query(Student).filter(Student.study_group_id != curator_group.id).first()
    task2 = _create(client, admin_headers, collect_mode="selected", fields=fields, title="Вторая")
    aid2 = _my_assignment_id(client, curator_headers, task2["id"])
    r = client.put(f"/tasks/assignments/{aid2}/answers", headers=curator_headers,
                   json={"rows": [{"student_id": other.id, "is_included": True}]})
    assert r.status_code == 400


# ---------- правила проверки и права ----------

def _submit_group_task(client, admin_headers, curator_headers, **overrides):
    fields = [{"label": "Сдано", "type": "bool", "required": True}]
    task = _create(client, admin_headers, collect_mode="group", fields=fields, **overrides)
    aid = _my_assignment_id(client, curator_headers, task["id"])
    client.put(f"/tasks/assignments/{aid}/answers", headers=curator_headers,
               json={"group_values": {task["fields"][0]["key"]: True}})
    return task, aid, client.post(f"/tasks/assignments/{aid}/submit", headers=curator_headers)


def test_reviewer_rule_edu_department(client, admin_headers, dept_head_headers, edu_department_headers, curator_headers, curator_group, imported):
    _, aid, r = _submit_group_task(client, admin_headers, curator_headers, reviewer_rule="edu_department")
    assert r.json()["status"] == "submitted"
    review = f"/tasks/assignments/{aid}/review"
    assert client.post(review, headers=dept_head_headers, json={"action": "accept"}).status_code == 403
    assert client.get("/tasks/review-queue", headers=dept_head_headers).json() == []
    assert len(client.get("/tasks/review-queue", headers=edu_department_headers).json()) == 1
    assert client.post(review, headers=edu_department_headers, json={"action": "accept"}).json()["status"] == "accepted"


def test_reviewer_rule_author_and_none(client, admin_headers, dept_head_headers, curator_headers, curator_group, imported):
    _, aid, r = _submit_group_task(client, dept_head_headers, curator_headers, reviewer_rule="author",
                                   title="От зав. отделением")
    assert r.json()["status"] == "submitted"
    assert client.post(f"/tasks/assignments/{aid}/review", headers=dept_head_headers,
                       json={"action": "accept"}).json()["status"] == "accepted"

    _, aid2, r2 = _submit_group_task(client, admin_headers, curator_headers, reviewer_rule="none", title="Без проверки")
    assert r2.json()["status"] == "accepted"


def test_access_is_limited_to_assignee_reviewer_and_scoped_managers(
    client, admin_headers, curator_headers, curator_group, db_second_curator, dept_head_headers, imported, db,
):
    task = _create(client, admin_headers)
    aid = _my_assignment_id(client, curator_headers, task["id"])
    from tests.conftest import _login

    other_headers = _login(client, db_second_curator.username, "SecondCurator123!")
    other_group_aid = next(
        a["id"] for a in task["assignments"] if a["study_group_id"] != curator_group.id
    )
    # Куратор не видит назначения чужой группы и не может заполнять чужое.
    assert client.get(f"/tasks/assignments/{other_group_aid}", headers=curator_headers).status_code == 403
    assert client.put(f"/tasks/assignments/{other_group_aid}/answers", headers=curator_headers,
                      json={"rows": []}).status_code == 403
    assert client.get(f"/tasks/assignments/{aid}", headers=other_headers).status_code == 403
    # Зав. отделением видит, но заполнять не может.
    assert client.get(f"/tasks/assignments/{aid}", headers=dept_head_headers).status_code == 200
    assert client.put(f"/tasks/assignments/{aid}/answers", headers=dept_head_headers,
                      json={"rows": []}).status_code == 403


def test_dept_head_of_another_department_sees_nothing(client, admin_headers, imported, db, dept_head_user):
    from app.core.security import hash_password
    from app.models import Role, User
    from tests.conftest import _login

    other = Department(name="Другое")
    db.add(other)
    db.flush()
    role = db.query(Role).filter(Role.code == "dept_head").one()
    db.add(User(username="other_dh", full_name="Чужой зав", role_id=role.id, department_id=other.id,
                password_hash=hash_password("Other123!")))
    db.commit()
    headers = _login(client, "other_dh", "Other123!")
    task = _create(client, admin_headers)
    assert client.get("/tasks", headers=headers).json() == []
    assert client.get(f"/tasks/{task['id']}", headers=headers).status_code == 403
    assert client.get(f"/tasks/{task['id']}/export", headers=headers).status_code == 403
    aid = task["assignments"][0]["id"]
    assert client.get(f"/tasks/assignments/{aid}", headers=headers).status_code == 403


def test_comments_notify_the_other_side(client, admin_headers, dept_head_headers, curator_headers, curator_group, imported, db):
    _, aid, _ = _submit_group_task(client, admin_headers, curator_headers)
    student = _students(db, curator_group)[0]
    r = client.post(f"/tasks/assignments/{aid}/comments", headers=dept_head_headers,
                    json={"student_id": student.id, "text": "Уточните данные"})
    assert r.status_code == 201
    assert any(n["kind"] == "task_comment" for n in client.get("/notifications", headers=curator_headers).json())
    comments = client.get(f"/tasks/assignments/{aid}", headers=curator_headers).json()["comments"]
    assert comments[0]["student_id"] == student.id and comments[0]["text"] == "Уточните данные"


# ---------- просрочка, управление, выгрузка ----------

def test_overdue_flag_and_progress(client, admin_headers, curator_headers, curator_group, imported, db):
    task = _create(client, admin_headers)
    db.query(Task).filter(Task.id == task["id"]).update({"due_date": datetime.date.today() - datetime.timedelta(days=2)})
    db.commit()
    detail = client.get(f"/tasks/{task['id']}", headers=admin_headers).json()
    assert detail["progress"]["overdue"] == detail["progress"]["total"]
    mine = client.get("/tasks/my", headers=curator_headers).json()
    assert mine[0]["is_overdue"] is True
    row = next(t for t in client.get("/tasks", headers=admin_headers).json() if t["id"] == task["id"])
    assert row["progress"]["overdue"] == row["progress"]["total"]


def test_close_update_and_delete_rules(client, admin_headers, dept_head_headers, curator_headers, curator_group, imported):
    task = _create(client, dept_head_headers)
    url = f"/tasks/{task['id']}"
    other = _create(client, admin_headers, title="Чужая")
    assert client.patch(f"/tasks/{other['id']}", headers=dept_head_headers, json={"title": "X"}).status_code == 403
    assert client.patch(url, headers=dept_head_headers, json={"title": "Новое имя"}).json()["title"] == "Новое имя"

    aid = _my_assignment_id(client, curator_headers, task["id"])
    assert client.patch(url, headers=dept_head_headers, json={"is_closed": True}).json()["is_closed"] is True
    r = client.put(f"/tasks/assignments/{aid}/answers", headers=curator_headers, json={"rows": []})
    assert r.status_code == 400 and "закрыта" in r.json()["detail"]
    client.patch(url, headers=dept_head_headers, json={"is_closed": False})
    assert client.delete(url, headers=admin_headers).status_code == 204  # ответов нет — можно удалить


def test_cannot_delete_task_with_answers(client, admin_headers, curator_headers, curator_group, imported):
    task, aid, _ = _submit_group_task(client, admin_headers, curator_headers)
    r = client.delete(f"/tasks/{task['id']}", headers=admin_headers)
    assert r.status_code == 400
    clean = _create(client, admin_headers, title="Без ответов")
    assert client.delete(f"/tasks/{clean['id']}", headers=admin_headers).status_code == 204
    assert client.get(f"/tasks/{clean['id']}", headers=admin_headers).status_code == 404


def test_export_per_student_and_per_group(client, admin_headers, curator_headers, curator_group, imported, db):
    task = _create(client, admin_headers)
    aid = _my_assignment_id(client, curator_headers, task["id"])
    keys = _field_keys(task)
    s = _students(db, curator_group)[0]
    client.put(f"/tasks/assignments/{aid}/answers", headers=curator_headers, json={"rows": [
        {"student_id": s.id, "values": {keys["Кружки"]: ["Спорт", "Танцы"], keys["Не посещает"]: False,
                                        keys["Телефон родителя"]: "+7 900"}}]})
    r = client.get(f"/tasks/{task['id']}/export", headers=admin_headers)
    assert r.status_code == 200
    rows = list(load_workbook(io.BytesIO(r.content)).worksheets[0].iter_rows(values_only=True))
    assert rows[0][:4] == ("Отделение", "Группа", "Статус", "Студент")
    assert len(rows) == 2  # только студент с ответом
    assert rows[1][1] == curator_group.code and rows[1][3] == s.full_name
    assert rows[1][4] == "Спорт, Танцы" and rows[1][5] == "нет" and rows[1][6] == "+7 900"

    group_task = _create(client, admin_headers, collect_mode="group",
                         fields=[{"label": "Сдано", "type": "bool"}], title="Групповая")
    rows = list(load_workbook(io.BytesIO(client.get(f"/tasks/{group_task['id']}/export", headers=admin_headers).content)).worksheets[0].iter_rows(values_only=True))
    assert rows[0][3] == "Сдано" and len(rows) == 1 + group_task["progress"]["total"]


# ---------- напоминания о сроках (проверка при открытии платформы) ----------

def _reminders(client, headers):
    return [n for n in client.get("/notifications", headers=headers).json()
            if n["kind"] in ("task_due_soon", "task_due_today", "task_overdue", "task_review_waiting")]


def _set_due(db, task_id, days):
    db.query(Task).filter(Task.id == task_id).update({"due_date": datetime.date.today() + datetime.timedelta(days=days)})
    db.commit()


def _open_platform(client, headers):
    """Как колокольчик при входе: опрос счётчика (и сброс 10-минутного ограничителя)."""
    from app.services import task_service

    task_service.reset_reminder_throttle()
    return client.get("/notifications/unread-count", headers=headers)


@pytest.mark.parametrize("days,kind", [(2, "task_due_soon"), (0, "task_due_today"), (-1, "task_overdue")])
def test_deadline_reminders_by_stage_and_only_once(client, admin_headers, curator_headers, curator_group, imported, db, days, kind):
    task = _create(client, admin_headers)
    _set_due(db, task["id"], days)
    assert _open_platform(client, curator_headers).status_code == 200
    first = _reminders(client, curator_headers)
    assert [n["kind"] for n in first] == [kind]
    assert first[0]["entity_type"] == "task_assignment"
    _open_platform(client, curator_headers)
    _open_platform(client, curator_headers)
    assert len(_reminders(client, curator_headers)) == 1  # повторно не создаётся


def test_no_reminders_for_far_closed_submitted_or_accepted(client, admin_headers, curator_headers, curator_group, imported, db):
    far = _create(client, admin_headers, title="Далёкая")
    closed = _create(client, admin_headers, title="Закрытая")
    done = _create(client, admin_headers, title="Сданная", collect_mode="group", fields=[{"label": "Сдано", "type": "bool", "required": True}])
    _set_due(db, far["id"], 10)
    _set_due(db, closed["id"], 1)
    client.patch(f"/tasks/{closed['id']}", headers=admin_headers, json={"is_closed": True})
    aid = _my_assignment_id(client, curator_headers, done["id"])
    client.put(f"/tasks/assignments/{aid}/answers", headers=curator_headers,
               json={"group_values": {done["fields"][0]["key"]: True}})
    client.post(f"/tasks/assignments/{aid}/submit", headers=curator_headers)
    _set_due(db, done["id"], 1)
    _open_platform(client, curator_headers)
    assert _reminders(client, curator_headers) == []


def test_reminders_are_throttled_between_polls(client, admin_headers, curator_headers, curator_group, imported, db):
    task = _create(client, admin_headers)
    _open_platform(client, curator_headers)  # первый запуск: срок далеко
    _set_due(db, task["id"], 1)
    client.get("/notifications/unread-count", headers=curator_headers)  # сразу же — в пределах 10 минут
    assert _reminders(client, curator_headers) == []
    _open_platform(client, curator_headers)  # «новое открытие» после паузы
    assert [n["kind"] for n in _reminders(client, curator_headers)] == ["task_due_soon"]


def _iso(days):
    return (datetime.date.today() + datetime.timedelta(days=days)).isoformat()


def test_extending_the_deadline_restarts_the_reminders(client, admin_headers, curator_headers, curator_group, imported, db):
    """Срок продлили после напоминания — раньше новых напоминаний не приходило вообще (каждое шлётся один раз)."""
    task = _create(client, admin_headers)
    _set_due(db, task["id"], -1)
    _open_platform(client, curator_headers)
    assert [n["kind"] for n in _reminders(client, curator_headers)] == ["task_overdue"]

    assert client.patch(f"/tasks/{task['id']}", headers=admin_headers, json={"due_date": _iso(10)}).status_code == 200
    assert _reminders(client, curator_headers) == []  # устаревшее «срок был …» убрано
    _open_platform(client, curator_headers)
    assert _reminders(client, curator_headers) == []  # срок далеко — тихо

    _set_due(db, task["id"], 2)  # подошёл новый срок
    _open_platform(client, curator_headers)
    assert [n["kind"] for n in _reminders(client, curator_headers)] == ["task_due_soon"]


def test_shortening_the_deadline_also_restarts_and_same_date_does_not(client, admin_headers, curator_headers, curator_group, imported, db):
    task = _create(client, admin_headers)
    _set_due(db, task["id"], 2)
    _open_platform(client, curator_headers)
    assert [n["kind"] for n in _reminders(client, curator_headers)] == ["task_due_soon"]

    # тот же срок (или правка названия) напоминания не трогает
    client.patch(f"/tasks/{task['id']}", headers=admin_headers, json={"due_date": _iso(2), "title": "Новое имя"})
    assert len(_reminders(client, curator_headers)) == 1

    client.patch(f"/tasks/{task['id']}", headers=admin_headers, json={"due_date": _iso(0)})
    _open_platform(client, curator_headers)
    assert [n["kind"] for n in _reminders(client, curator_headers)] == ["task_due_today"]


def test_changing_the_deadline_keeps_review_waiting_reminders(client, admin_headers, dept_head_headers, curator_headers, curator_group, imported, db):
    task, aid, r = _submit_group_task(client, admin_headers, curator_headers)
    db.query(TaskAssignment).filter(TaskAssignment.id == aid).update({"submitted_at": utcnow() - datetime.timedelta(days=3)})
    db.commit()
    _open_platform(client, dept_head_headers)
    assert [n["kind"] for n in _reminders(client, dept_head_headers)] == ["task_review_waiting"]
    client.patch(f"/tasks/{task['id']}", headers=admin_headers, json={"due_date": _iso(20)})
    assert [n["kind"] for n in _reminders(client, dept_head_headers)] == ["task_review_waiting"]


def test_reviewer_is_reminded_about_stale_submissions(client, admin_headers, dept_head_headers, curator_headers, curator_group, imported, db):
    task, aid, r = _submit_group_task(client, admin_headers, curator_headers)
    assert r.status_code == 200
    _open_platform(client, dept_head_headers)
    assert _reminders(client, dept_head_headers) == []  # только что сдано
    db.query(TaskAssignment).filter(TaskAssignment.id == aid).update(
        {"submitted_at": utcnow() - datetime.timedelta(days=3)})
    db.commit()
    _open_platform(client, dept_head_headers)
    _open_platform(client, dept_head_headers)
    assert [n["kind"] for n in _reminders(client, dept_head_headers)] == ["task_review_waiting"]
    # Кто не проверяет эту задачу, напоминания не получает.
    _open_platform(client, curator_headers)
    assert all(n["kind"] != "task_review_waiting" for n in _reminders(client, curator_headers))


# ---------- двухступенчатая проверка и новые группы ----------

def test_two_step_review_dept_head_then_edu_department(
    client, admin_headers, dept_head_headers, edu_department_headers, curator_headers, curator_group, imported,
):
    _, aid, r = _submit_group_task(client, admin_headers, curator_headers, reviewer_rule="two_step")
    assert r.json()["status"] == "submitted" and r.json()["review_step"] == 1 and r.json()["review_steps"] == 2
    url = f"/tasks/assignments/{aid}/review"

    # Сначала проверяет зав. отделением; воспитательный отдел пока не может.
    assert client.post(url, headers=edu_department_headers, json={"action": "accept"}).status_code == 403
    assert client.get("/tasks/review-queue", headers=edu_department_headers).json() == []
    assert len(client.get("/tasks/review-queue", headers=dept_head_headers).json()) == 1

    first = client.post(url, headers=dept_head_headers, json={"action": "accept"}).json()
    assert first["status"] == "submitted" and first["review_step"] == 2 and first["can_review"] is False
    assert any(n["kind"] == "task_submitted" for n in client.get("/notifications", headers=edu_department_headers).json())
    assert client.get("/tasks/review-queue", headers=dept_head_headers).json() == []
    assert client.post(url, headers=dept_head_headers, json={"action": "accept"}).status_code == 403
    # Куратор об «принято» ещё не уведомлён и править не может.
    assert not any(n["kind"] == "task_accepted" for n in client.get("/notifications", headers=curator_headers).json())
    assert client.get(f"/tasks/assignments/{aid}", headers=curator_headers).json()["can_edit"] is False

    final = client.post(url, headers=edu_department_headers, json={"action": "accept"}).json()
    assert final["status"] == "accepted"
    assert any(n["kind"] == "task_accepted" for n in client.get("/notifications", headers=curator_headers).json())


def test_two_step_return_at_second_step_restarts_from_first(
    client, admin_headers, dept_head_headers, edu_department_headers, curator_headers, curator_group, imported,
):
    task, aid, _ = _submit_group_task(client, admin_headers, curator_headers, reviewer_rule="two_step")
    url = f"/tasks/assignments/{aid}/review"
    client.post(url, headers=dept_head_headers, json={"action": "accept"})
    r = client.post(url, headers=edu_department_headers, json={"action": "return", "comment": "Не хватает данных"})
    assert r.json()["status"] == "returned" and r.json()["review_step"] == 1
    client.put(f"/tasks/assignments/{aid}/answers", headers=curator_headers,
               json={"group_values": {task["fields"][0]["key"]: True}})
    assert client.post(f"/tasks/assignments/{aid}/submit", headers=curator_headers).json()["review_step"] == 1
    assert len(client.get("/tasks/review-queue", headers=dept_head_headers).json()) == 1  # снова у зав. отделением


def test_new_group_gets_open_tasks_in_scope(client, admin_headers, dept_head_headers, dept_head_user, imported, db):
    everything = _create(client, admin_headers)
    only_course_3 = _create(client, admin_headers, title="Для 3 курса", scope={"courses": [3]})
    explicit = _create(client, admin_headers, title="Выбранные", scope={"group_ids": [db.query(StudyGroup).first().id]})
    by_dept_head = _create(client, dept_head_headers, title="От зав. отделением")
    closed = _create(client, admin_headers, title="Закрытая")
    client.patch(f"/tasks/{closed['id']}", headers=admin_headers, json={"is_closed": True})
    expired = _create(client, admin_headers, title="Истёкшая")
    db.query(Task).filter(Task.id == expired["id"]).update({"due_date": datetime.date.today() - datetime.timedelta(days=1)})
    db.commit()

    r = client.post("/admin/groups", headers=admin_headers,
                    json={"code": "NEW-1", "course": 1, "department_id": dept_head_user.department_id})
    assert r.status_code == 201
    new_id = r.json()["id"]
    got = {a.task_id for a in db.query(TaskAssignment).filter(TaskAssignment.study_group_id == new_id)}
    assert got == {everything["id"], by_dept_head["id"]}  # курс 3, явные группы, закрытая и истёкшая — нет

    # Группа из чужого отделения не получает задачу зав. отделением.
    other = Department(name="Другое")
    db.add(other)
    db.commit()
    r = client.post("/admin/groups", headers=admin_headers, json={"code": "NEW-2", "course": 1, "department_id": other.id})
    got = {a.task_id for a in db.query(TaskAssignment).filter(TaskAssignment.study_group_id == r.json()["id"])}
    assert got == {everything["id"]}
    assert only_course_3["id"] not in got and explicit["id"] not in got


def test_restored_group_gets_missing_assignments(client, admin_headers, imported, db):
    group = db.query(StudyGroup).filter(StudyGroup.is_active.is_(True)).first()
    client.patch(f"/admin/groups/{group.id}", headers=admin_headers, json={"is_active": False})
    task = _create(client, admin_headers)  # архивной группе назначения не создаются
    assert db.query(TaskAssignment).filter(TaskAssignment.task_id == task["id"], TaskAssignment.study_group_id == group.id).count() == 0
    client.patch(f"/admin/groups/{group.id}", headers=admin_headers, json={"is_active": True})
    db.expire_all()
    assert db.query(TaskAssignment).filter(TaskAssignment.task_id == task["id"], TaskAssignment.study_group_id == group.id).count() == 1
    client.patch(f"/admin/groups/{group.id}", headers=admin_headers, json={"is_active": False})
    client.patch(f"/admin/groups/{group.id}", headers=admin_headers, json={"is_active": True})  # повторно — без дублей
    assert db.query(TaskAssignment).filter(TaskAssignment.task_id == task["id"], TaskAssignment.study_group_id == group.id).count() == 1


def test_export_neutralizes_formulas_typed_into_answers(client, admin_headers, curator_headers, curator_group, imported, db):
    """Ответ куратора «=…» не должен превратиться в формулу у того, кто откроет Excel."""
    task = _create(client, admin_headers, collect_mode="group", title="=1+1",
                   fields=[{"label": "=HYPERLINK(\"http://evil\")", "type": "text"}])
    aid = _my_assignment_id(client, curator_headers, task["id"])
    client.put(f"/tasks/assignments/{aid}/answers", headers=curator_headers,
               json={"group_values": {task["fields"][0]["key"]: "=cmd|' /C calc'!A0"}})
    ws = load_workbook(io.BytesIO(client.get(f"/tasks/{task['id']}/export", headers=admin_headers).content)).worksheets[0]
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any(c.value.startswith("=cmd") for c in cells)  # содержимое не искажено…
    assert all(c.data_type != "f" for c in cells)  # …но формулой не становится


# ---------- кто и когда проверил ----------

def _history(client, headers, aid):
    d = client.get(f"/tasks/assignments/{aid}", headers=headers).json()
    return [(e["kind"], e["user_name"], e["step"]) for e in d["history"]], d


def test_history_shows_who_submitted_returned_and_accepted(
    client, admin_headers, dept_head_headers, dept_head_user, curator_headers, curator_user, curator_group, imported,
):
    task, aid, _ = _submit_group_task(client, admin_headers, curator_headers)
    key = task["fields"][0]["key"]
    events, d = _history(client, curator_headers, aid)
    assert events == [("submitted", curator_user.full_name, None)]
    assert d["submitted_at"] and d["reviewed_at"] is None and d["reviewed_by_name"] is None

    client.post(f"/tasks/assignments/{aid}/review", headers=dept_head_headers, json={"action": "return", "comment": "Дополните"})
    client.put(f"/tasks/assignments/{aid}/answers", headers=curator_headers, json={"group_values": {key: True}})
    client.post(f"/tasks/assignments/{aid}/submit", headers=curator_headers)
    client.post(f"/tasks/assignments/{aid}/review", headers=dept_head_headers, json={"action": "accept"})

    events, d = _history(client, dept_head_headers, aid)
    assert events == [
        ("submitted", curator_user.full_name, None), ("returned", dept_head_user.full_name, None),
        ("submitted", curator_user.full_name, None), ("accepted", dept_head_user.full_name, None),
    ]
    assert d["reviewed_by_name"] == dept_head_user.full_name and d["reviewed_at"] is not None


def test_two_step_history_labels_each_step(
    client, admin_headers, dept_head_headers, dept_head_user, edu_department_headers, edu_department_user,
    curator_headers, curator_group, imported,
):
    _, aid, _ = _submit_group_task(client, admin_headers, curator_headers, reviewer_rule="two_step")
    url = f"/tasks/assignments/{aid}/review"
    client.post(url, headers=dept_head_headers, json={"action": "accept"})
    client.post(url, headers=edu_department_headers, json={"action": "accept"})
    events, d = _history(client, admin_headers, aid)
    assert [(k, s) for k, _, s in events] == [("submitted", None), ("accepted", 1), ("accepted", 2)]
    assert [n for _, n, _ in events][1:] == [dept_head_user.full_name, edu_department_user.full_name]
    assert d["reviewed_by_name"] == edu_department_user.full_name  # итоговое решение — за второй ступенью


def test_accepted_without_review_is_marked_as_automatic(client, admin_headers, curator_headers, curator_group, imported):
    _, aid, _ = _submit_group_task(client, admin_headers, curator_headers, reviewer_rule="none")
    events, d = _history(client, curator_headers, aid)
    assert [(k, n) for k, n, _ in events][-1] == ("auto_accepted", None)
    assert d["status"] == "accepted" and d["reviewed_by_name"] is None


def test_matrix_shows_reviewer_names_and_history_is_per_assignment(
    client, admin_headers, dept_head_headers, dept_head_user, curator_headers, curator_group, imported,
):
    task, aid, _ = _submit_group_task(client, admin_headers, curator_headers)
    client.post(f"/tasks/assignments/{aid}/review", headers=dept_head_headers, json={"action": "accept"})
    detail = client.get(f"/tasks/{task['id']}", headers=admin_headers).json()
    mine = next(a for a in detail["assignments"] if a["id"] == aid)
    others = [a for a in detail["assignments"] if a["id"] != aid]
    assert mine["reviewed_by_name"] == dept_head_user.full_name and mine["reviewed_at"]
    assert all(a["reviewed_by_name"] is None for a in others)
    # история чужого назначения не подмешивается
    other_events = client.get(f"/tasks/assignments/{others[0]['id']}", headers=admin_headers).json()["history"]
    assert other_events == []
