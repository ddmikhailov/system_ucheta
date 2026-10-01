"""«К какой паре пришли» (DaySubmission.first_period) и свод посещаемости
«всего / к 1 паре» в выгрузке (summary_service + листы в export_service)."""
import datetime
import io

import openpyxl
import pytest

from app.models import Role, RoleCode, StudyGroup, User
from app.services import attendance_service, calendar_service, summary_service


def _recent_study_day(db, group, today):
    day = today
    for _ in range(14):
        if calendar_service.is_study_day(db, day, study_group_id=group.id, course=group.course):
            return day
        day -= datetime.timedelta(days=1)
    raise AssertionError("нет учебного дня за две недели")


def _admin(db):
    return db.query(User).join(Role).filter(Role.code == RoleCode.ADMIN.value).one()


def test_submit_day_saves_first_period_and_resubmit_keeps_it(client, curator_headers, curator_group, db, today):
    day = _recent_study_day(db, curator_group, today)

    r = client.post(
        f"/curator/groups/{curator_group.id}/day/submit?date={day}",
        headers=curator_headers, json={"exceptions": [], "first_period": 2},
    )
    assert r.status_code == 200, r.text
    assert r.json()["first_period"] == 2

    # «Все присутствуют» без выбора пары не стирает уже выставленную.
    r = client.post(
        f"/curator/groups/{curator_group.id}/day/mark-all-present?date={day}&confirm=true",
        headers=curator_headers,
    )
    assert r.status_code == 200, r.text
    assert r.json()["first_period"] == 2

    r = client.post(
        f"/curator/groups/{curator_group.id}/day/mark-all-present?date={day}&confirm=true&first_period=1",
        headers=curator_headers,
    )
    assert r.json()["first_period"] == 1

    r = client.get(f"/curator/groups/{curator_group.id}/day?date={day}", headers=curator_headers)
    assert r.json()["first_period"] == 1


def test_first_period_out_of_range_rejected(client, curator_headers, curator_group, db, today):
    day = _recent_study_day(db, curator_group, today)
    r = client.post(
        f"/curator/groups/{curator_group.id}/day/submit?date={day}",
        headers=curator_headers, json={"exceptions": [], "first_period": 11},
    )
    assert r.status_code == 422


def test_submit_day_service_rejects_invalid_period(db, curator_group, today):
    day = _recent_study_day(db, curator_group, today)
    with pytest.raises(attendance_service.InvalidSubmission):
        attendance_service.submit_day(db, curator_group.id, day, [], _admin(db), first_period=0)


def _two_groups_with_day(db, today):
    groups = db.query(StudyGroup).filter(StudyGroup.course == 2, StudyGroup.is_active.is_(True)).limit(3).all()
    assert len(groups) == 3
    day = _recent_study_day(db, groups[0], today)
    for g in groups:
        assert calendar_service.is_study_day(db, day, study_group_id=g.id, course=g.course)
    return groups, day


def test_summary_splits_total_and_first_period(db, imported, today):
    (first, second, unsubmitted), day = _two_groups_with_day(db, today)
    admin = _admin(db)

    students = attendance_service.get_active_students(db, first.id, day)
    attendance_service.submit_day(
        db, first.id, day,
        [{"student_id": students[0].id, "mark_code": "н", "comment": None, "basis_reference": None}],
        admin, first_period=1,
    )
    attendance_service.submit_day(db, second.id, day, [], admin, first_period=2)

    rows = summary_service.collect_group_day_rows(db, day, day)
    by_code = {r.group_code: r for r in rows}
    assert len(rows) == len([g for g in db.query(StudyGroup).filter(StudyGroup.is_active.is_(True))
                              if calendar_service.is_study_day(db, day, study_group_id=g.id, course=g.course)])

    a, b, c = by_code[first.code], by_code[second.code], by_code[unsubmitted.code]
    assert (a.is_submitted, a.first_period, a.absent, a.by_code) == (True, 1, 1, {"н": 1})
    assert a.present == a.headcount - 1
    assert (b.is_submitted, b.first_period, b.absent) == (True, 2, 0)
    assert c.is_submitted is False and c.present == 0

    everyone = summary_service.totals_for(rows)
    assert everyone.headcount == sum(r.headcount for r in rows)
    assert everyone.counted == a.headcount + b.headcount  # несданная группа не «все пришли»
    assert everyone.absent == 1
    assert everyone.by_code == {"н": 1}

    first_only = summary_service.totals_for([r for r in rows if summary_service.is_first_period(r)])
    assert first_only.groups_submitted == 1
    assert first_only.counted == a.headcount
    assert first_only.present == a.headcount - 1
    assert first_only.percent == round((a.headcount - 1) / a.headcount * 100, 2)


def test_excel_export_has_summary_sheets(client, admin_headers, db, imported, today):
    (first, second, _), day = _two_groups_with_day(db, today)
    admin = _admin(db)
    students = attendance_service.get_active_students(db, first.id, day)
    attendance_service.submit_day(
        db, first.id, day,
        [{"student_id": students[0].id, "mark_code": "б", "comment": None, "basis_reference": None}],
        admin, first_period=1,
    )
    attendance_service.submit_day(db, second.id, day, [], admin, first_period=3)

    r = client.get(f"/export/excel?date_from={day}&date_to={day}", headers=admin_headers)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:4] == ["ИТОГ", "Свод по дням", "Свод за период", "По группам и дням"]

    daily = list(wb["Свод по дням"].iter_rows(values_only=True))
    header = daily[0]
    assert header[:3] == ("Дата", "Отделение", "Срез")
    body = {(row[1], row[2]): row for row in daily[1:]}
    total = body[("Диджитал", "Всего")]
    first_period = body[("Диджитал", "К 1 паре")]
    col = {name: i for i, name in enumerate(header)}
    assert first_period[col["Групп сдали"]] == 1
    assert first_period[col["Учтено студентов"]] == len(students)
    assert first_period[col["Отсутствуют"]] == 1
    assert total[col["Групп сдали"]] == 2
    assert total[col["Учтено студентов"]] > first_period[col["Учтено студентов"]]
    sick_column = next(i for i, name in enumerate(header) if str(name).startswith("Б — "))
    assert total[sick_column] == 1
    assert first_period[sick_column] == 1

    per_group = list(wb["По группам и дням"].iter_rows(values_only=True))
    pg_header = per_group[0]
    rows = {row[2]: row for row in per_group[1:]}
    assert rows[first.code][pg_header.index("К какой паре пришли")] == 1
    assert rows[second.code][pg_header.index("К какой паре пришли")] == 3


def test_dashboards_summary_endpoint_matches_service_numbers(client, admin_headers, db, imported, today):
    (first, second, _), day = _two_groups_with_day(db, today)
    admin = _admin(db)
    students = attendance_service.get_active_students(db, first.id, day)
    attendance_service.submit_day(
        db, first.id, day,
        [{"student_id": students[0].id, "mark_code": "н", "comment": None, "basis_reference": None}],
        admin, first_period=1,
    )
    attendance_service.submit_day(db, second.id, day, [], admin, first_period=2)

    r = client.get(f"/dashboards/summary?date_from={day}&date_to={day}", headers=admin_headers)
    assert r.status_code == 200, r.text
    body = r.json()
    assert [c["code"] for c in body["codes"]][:2] == ["о", "и"]
    assert body["group_days"] == []  # тяжёлая разбивка — только по запросу

    lines = {(x["department"], x["slice_name"]): x for x in body["daily"]}
    total, first_period = lines[("Диджитал", "Всего")], lines[("Диджитал", "К 1 паре")]
    assert first_period["groups_submitted"] == 1
    assert first_period["counted"] == len(students)
    assert first_period["absent"] == 1
    assert first_period["by_code"] == {"н": 1}
    assert total["groups_submitted"] == 2
    assert total["counted"] > first_period["counted"]
    assert total["date"] == str(day)

    period = {(x["department"], x["slice_name"]): x for x in body["period"]}
    assert period[("Диджитал", "К 1 паре")]["date"] is None
    assert period[("Диджитал", "К 1 паре")]["counted"] == first_period["counted"]

    r = client.get(
        f"/dashboards/summary?date_from={day}&date_to={day}&include_group_days=true", headers=admin_headers
    )
    by_group = {g["group_code"]: g for g in r.json()["group_days"]}
    assert by_group[first.code]["first_period"] == 1
    assert by_group[first.code]["absent"] == 1
    assert by_group[second.code]["first_period"] == 2
    unsubmitted = next(g for g in by_group.values() if not g["is_submitted"])
    assert unsubmitted["present"] is None and unsubmitted["percent"] is None


def test_dashboards_summary_scoped_and_protected(client, dept_head_headers, curator_headers, db, imported, today):
    group = db.query(StudyGroup).filter(StudyGroup.course == 2).first()
    day = _recent_study_day(db, group, today)
    r = client.get(f"/dashboards/summary?date_from={day}&date_to={day}", headers=dept_head_headers)
    assert r.status_code == 200
    assert {x["department"] for x in r.json()["daily"]} <= {"Диджитал"}
    assert client.get(f"/dashboards/summary?date_from={day}&date_to={day}", headers=curator_headers).status_code == 403
